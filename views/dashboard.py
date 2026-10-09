from datetime import datetime
from hashlib import sha256
from html import escape
from html.parser import HTMLParser
import io
import json
import math
from pathlib import Path
from urllib.parse import quote
import zipfile
from xml.etree import ElementTree

import pandas as pd
import plotly.express as px
import plotly.graph_objects as go
from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException
from PIL import Image, UnidentifiedImageError
from pypdf import PdfReader
from pypdf.errors import PdfReadError
import streamlit as st

from api_client import (
    APIError,
    AuthenticationError,
    create_quote,
    get_building_quote,
    get_buildings,
    get_formula,
    get_metrics,
    get_model_health,
    get_portfolio_summary,
    model_get,
    upload_document,
)
from interpretation_client import (
    ask_uploaded_document_question,
    ask_interpretation_question,
    check_interpretation_health,
    get_interpretation_base_url,
    request_uploaded_document_summary,
    request_interpretation_summary,
)


TOKEN = st.session_state.get("token")
UPLOAD_EXTENSIONS = (
    "txt", "md", "docx", "pdf", "xlsx", "csv", "html", "json", "py",
    "png", "jpg", "jpeg", "svg", "tif", "tiff", "zip",
)
MAX_UPLOAD_BYTES = 25 * 1024 * 1024
MAX_UPLOAD_TOTAL_BYTES = 50 * 1024 * 1024
MAX_ARCHIVE_ENTRIES = 100
MAX_ARCHIVE_EXPANDED_BYTES = 25 * 1024 * 1024
MAX_REPORT_TEXT_CHARS = 6000
TEAL = "#006d83"
TEAL_DARK = "#073744"
AQUA = "#1b9a9e"
BG = "#eff8fb"
LINE = "#cfe0e6"
TEXT = "#00243d"
MUTED = "#476074"
CLASS_COLORS = ["#dc4b26", "#e89b2c", "#279d95", "#075e6e", "#7fb3bf"]
BROKER_PLACEMENT = {
    "reference": "EIB-NAI-LP-2026-001",
    "client": "Landmark Plaza Commercial Development",
    "location": "Upper Hill, Nairobi",
    "issued": "8 Oct 2026",
    "expiry": "12 Oct 2026",
    "insured_value_kes": 1_090_000_000,
    "current_premium_kes": 2_450_000,
    "historical_losses_kes": 1_210_000,
    "historical_loss_years": 11,
    "coverage": "All-risks excluding flood; flood cover available on request",
    "flood_limit": "Full TIV (KES 1.09B)",
    "flood_deductible": "5% or KES 5M minimum; basis to confirm",
    "reported_flood_losses": "None reported since construction in 2015",
    "broker_risk_profile": "A-class (broker assessment)",
    "attention_items": [
        "Basement 2 underground fuel tank: confirm environmental safeguards.",
        "Roof membrane is 11 years old; replacement planning noted for 2027–2028.",
        "2 of 48 emergency lights were reported non-functional.",
        "Monitor minor, reportedly non-structural basement wall cracks.",
    ],
}


def inject_styles():
    st.markdown(
        """
        <style>
        @import url('https://fonts.googleapis.com/css2?family=Inter:wght@400;600;700;800&display=swap');

        html, body, [class*="css"] {
            font-family: Inter, Arial, sans-serif;
        }

        .stApp {
            background: #eff8fb;
            color: #00243d;
        }

        header[data-testid="stHeader"] {
            background: transparent;
        }

        section[data-testid="stSidebar"] {
            background: #073744;
            border-right: 1px solid rgba(255,255,255,.08);
        }

        section[data-testid="stSidebar"] * {
            color: rgba(255,255,255,.72);
        }

        section[data-testid="stSidebar"] h1,
        section[data-testid="stSidebar"] h2,
        section[data-testid="stSidebar"] h3,
        section[data-testid="stSidebar"] strong {
            color: #fff;
        }

        section[data-testid="stSidebar"] div[role="radiogroup"] label {
            padding: 12px 14px;
            border-radius: 14px;
            margin: 6px 0;
            font-weight: 700;
        }

        section[data-testid="stSidebar"] div[role="radiogroup"] label:hover {
            background: rgba(23, 154, 158, .18);
        }

        section[data-testid="stSidebar"] div[role="radiogroup"] label[data-baseweb="radio"] > div:first-child {
            display: none;
        }

        .block-container {
            max-width: 1580px;
            padding-top: 2rem;
            padding-bottom: 4rem;
        }

        h1 {
            color: #00243d;
            font-size: 2.35rem !important;
            line-height: 1.05 !important;
            letter-spacing: 0 !important;
            font-weight: 800 !important;
        }

        h2, h3 {
            color: #00243d;
            letter-spacing: 0 !important;
        }

        .hero-subtitle {
            color: #476074;
            font-size: 1.12rem;
            line-height: 1.45;
            margin-top: -8px;
            max-width: 920px;
        }

        .welcome-date {
            color: #728493;
            font-size: .95rem;
            font-weight: 600;
            margin: 0 0 9px;
        }

        .welcome-greeting {
            color: #102d40;
            font-size: 2.25rem;
            font-weight: 750;
            letter-spacing: -.045em;
            line-height: 1.15;
            margin: 0;
        }

        .welcome-subtitle {
            color: #61788a;
            font-size: 1rem;
            margin-top: 9px;
        }

        div.st-key-broker_notification_popover [data-testid="stPopover"] > button {
            position: relative;
            width: 48px;
            min-height: 48px;
            padding: 0;
            border: 1px solid #006174;
            border-radius: 50%;
            color: #fff;
            background: #00758c;
            box-shadow: 0 4px 12px rgba(0, 77, 95, .22);
            transition: background .18s ease, box-shadow .18s ease, transform .18s ease;
        }

        div.st-key-broker_notification_popover [data-testid="stPopover"] > button:hover {
            color: #fff;
            border-color: #004f60;
            background: #00586a;
            box-shadow: 0 6px 16px rgba(0, 77, 95, .28);
            transform: translateY(-1px);
        }

        div.st-key-broker_notification_popover [data-testid="stPopover"] > button:focus-visible {
            outline: 3px solid #f2b544;
            outline-offset: 3px;
        }

        div.st-key-broker_notification_popover [data-testid="stPopover"] > button p {
            display: none;
        }

        .broker-message-preview {
            padding: 12px 14px;
            border: 1px solid #e4ebef;
            border-radius: 13px;
            background: #f8fbfc;
        }

        .broker-message-preview p {
            margin: 0;
            line-height: 1.5;
        }

        .nf-card {
            background: white;
            border: 1px solid #cfe0e6;
            border-radius: 18px;
            box-shadow: 0 10px 24px rgba(7, 55, 68, .08);
            padding: 24px;
            min-height: 100%;
        }

        .nf-metric {
            background: white;
            border: 1px solid #cfe0e6;
            border-radius: 18px;
            box-shadow: 0 10px 24px rgba(7, 55, 68, .08);
            padding: 24px 26px;
            min-height: 172px;
        }

        .nf-metric-icon {
            color: #006d83;
            font-size: 24px;
            font-weight: 800;
            margin-bottom: 28px;
        }

        .nf-metric-label {
            color: #476074;
            font-size: .98rem;
            margin-bottom: 6px;
            text-decoration: underline;
            text-decoration-color: #9ab4bf;
            text-underline-offset: 3px;
        }

        .nf-metric-value {
            color: #00243d;
            font-size: 1.9rem;
            line-height: 1.1;
            font-weight: 800;
        }

        .nf-metric-help {
            color: #476074;
            font-size: .88rem;
            margin-top: 8px;
        }

        .badge-row {
            display: flex;
            gap: 8px;
            flex-wrap: wrap;
            justify-content: flex-end;
            margin-top: 8px;
        }

        .badge {
            display: inline-flex;
            align-items: center;
            width: fit-content;
            padding: 4px 9px;
            border-radius: 999px;
            border: 1px solid #b9cdd6;
            background: #edf4f7;
            color: #476074;
            font-family: monospace;
            font-size: .74rem;
            letter-spacing: .02em;
            text-transform: uppercase;
        }

        .badge.source-icon {
            width: 28px;
            height: 28px;
            justify-content: center;
            padding: 0;
            font-family: Inter, Arial, sans-serif;
            font-size: 1rem;
            text-transform: none;
        }

        .badge.synthetic {
            border-color: #9bc9ff;
            color: #1f5b9d;
            background: #edf6ff;
        }

        .badge.real {
            border-color: #95d7bb;
            color: #00693f;
            background: #e9f8f0;
        }

        .badge.assumption {
            border-color: #e9b261;
            color: #a15c00;
            background: #fff2df;
        }

        .story-card {
            font-size: 1rem;
            line-height: 1.55;
        }

        .story-card b {
            color: #00243d;
        }

        .notice {
            border-radius: 14px;
            padding: 14px 16px;
            border: 1px solid #95d7de;
            background: #dff7f7;
            color: #00384d;
            margin: 14px 0 20px;
        }

        .warn-note {
            border-radius: 14px;
            padding: 14px 16px;
            border: 1px solid #f0be77;
            background: #fff5e6;
            color: #573300;
            margin: 14px 0 20px;
        }

        div[data-testid="stDataFrame"] {
            border: 1px solid #cfe0e6;
            border-radius: 14px;
            overflow: hidden;
            background: white;
        }

        .stButton > button,
        .stDownloadButton > button {
            border-radius: 18px;
            border: 1px solid #00758c;
            background: linear-gradient(135deg, #00758c, #1a9c9e);
            color: white;
            font-weight: 800;
            min-height: 44px;
        }

        .stButton > button:hover,
        .stDownloadButton > button:hover {
            border-color: #00586a;
            color: white;
        }

        div.st-key-explore_ep_curve button,
        div.st-key-explore_hotspots button,
        div.st-key-explore_find_quote button,
        div.st-key-explore_methodology button {
            width: 100%;
            min-height: 48px;
            justify-content: space-between;
            border: 1px solid #cfe0e6;
            border-radius: 14px;
            background: #fff;
            color: #00243d;
            box-shadow: none;
            font-weight: 600;
            text-align: left;
        }

        div.st-key-explore_ep_curve button:hover,
        div.st-key-explore_hotspots button:hover,
        div.st-key-explore_find_quote button:hover,
        div.st-key-explore_methodology button:hover {
            border-color: #82b9c5;
            background: #f6fbfc;
            color: #005b70;
        }

        div.st-key-header_generate_report button {
            min-height: 48px;
            border-color: #15803d;
            border-radius: 14px;
            background: #15803d;
            color: #fff;
            box-shadow: 0 4px 12px rgba(21, 128, 61, .2);
            white-space: nowrap;
        }

        div.st-key-header_generate_report button:hover {
            border-color: #166534;
            background: #166534;
            color: #fff;
            box-shadow: 0 6px 16px rgba(21, 128, 61, .28);
        }

        div.st-key-header_generate_report button:focus-visible {
            outline: 3px solid #f2b544;
            outline-offset: 3px;
        }

        div[data-testid="stTextInput"] input,
        div[data-testid="stSelectbox"] div[data-baseweb="select"] > div,
        div[data-testid="stMultiSelect"] div[data-baseweb="select"] > div {
            border-radius: 14px;
            border-color: #cfe0e6;
            background: #f4fbfd;
        }
        </style>
        """,
        unsafe_allow_html=True,
    )


def prettify(label):
    return str(label).replace("_", " ").replace("-", " ").title()


def is_number(value):
    return isinstance(value, (int, float)) and not isinstance(value, bool)


def to_number(value):
    if is_number(value):
        return float(value)
    if isinstance(value, str):
        cleaned = value.replace(",", "").replace("KES", "").replace("%", "").strip()
        try:
            return float(cleaned)
        except ValueError:
            return None
    return None


def format_number(value):
    number = to_number(value)
    if number is None:
        return "-"
    if abs(number) >= 1_000_000_000:
        return f"{number / 1_000_000_000:.2f}B"
    if abs(number) >= 1_000_000:
        return f"{number / 1_000_000:.0f}M"
    if abs(number) >= 1_000:
        return f"{number / 1_000:.1f}K"
    return f"{number:,.0f}"


def format_kes(value):
    number = to_number(value)
    if number is None:
        return "-"
    return f"KES {format_number(number)}"


def format_pct(value):
    number = to_number(value)
    if number is None:
        return "-"
    if number > 1:
        return f"{number:.1f}%"
    return f"{number * 100:.1f}%"


def safe(value):
    return escape(str(value)) if value is not None else ""


def normalize_records(payload):
    if isinstance(payload, list):
        return payload

    if isinstance(payload, dict):
        for key in ("results", "buildings", "data", "items", "records"):
            value = payload.get(key)
            if isinstance(value, list):
                return value

    return []


class HTMLTextExtractor(HTMLParser):
    def __init__(self):
        super().__init__(convert_charrefs=True)
        self.parts = []
        self._ignored_tag_depth = 0

    def handle_starttag(self, tag, _):
        if tag.casefold() in {"script", "style"}:
            self._ignored_tag_depth += 1

    def handle_endtag(self, tag):
        if tag.casefold() in {"script", "style"} and self._ignored_tag_depth:
            self._ignored_tag_depth -= 1

    def handle_data(self, data):
        if not self._ignored_tag_depth and data.strip():
            self.parts.append(data.strip())


def decode_uploaded_text(content):
    return content.decode("utf-8-sig", errors="replace").strip()


def summarize_csv(content):
    frame = pd.read_csv(io.BytesIO(content))
    preview = frame.head(8).to_string(index=False)
    return (
        f"Rows: {len(frame):,}\nColumns ({len(frame.columns)}): "
        f"{', '.join(str(column) for column in frame.columns)}\n"
        f"First rows:\n{preview}"
    )


def summarize_workbook(content):
    workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    sheets = []
    try:
        for sheet in workbook.worksheets:
            preview_rows = [
                ["" if value is None else str(value) for value in row]
                for row in sheet.iter_rows(
                    min_row=1, max_row=min(sheet.max_row or 1, 8),
                    max_col=min(sheet.max_column or 1, 12), values_only=True,
                )
            ]
            preview = "\n".join(" | ".join(row) for row in preview_rows)
            sheets.append(
                f"Sheet: {sheet.title} ({sheet.max_row} rows, "
                f"{sheet.max_column} columns)\n{preview}"
            )
    finally:
        workbook.close()
    return "\n\n".join(sheets) or "Workbook contains no sheets."


def summarize_docx(content):
    with zipfile.ZipFile(io.BytesIO(content)) as archive:
        document = ElementTree.fromstring(archive.read("word/document.xml"))
    paragraphs = []
    for paragraph in document.iter():
        if paragraph.tag.rsplit("}", 1)[-1] == "p":
            text = "".join(
                node.text or ""
                for node in paragraph.iter()
                if node.tag.rsplit("}", 1)[-1] == "t"
            ).strip()
            if text:
                paragraphs.append(text)
    return "\n".join(paragraphs) or "No readable text was found in the document."


def summarize_pdf(content):
    reader = PdfReader(io.BytesIO(content), strict=False)
    if reader.is_encrypted:
        raise ValueError("This PDF is password-protected and cannot be read.")
    page_text = []
    for page_number, page in enumerate(reader.pages[:40], start=1):
        text = (page.extract_text() or "").strip()
        if text:
            page_text.append(f"Page {page_number}: {text}")
    result = f"Pages: {len(reader.pages)}\n" + "\n\n".join(page_text)
    if not page_text:
        result += "\nNo embedded text found (the PDF may be scanned)."
    elif len(reader.pages) > 40:
        result += "\nOnly the first 40 pages were included."
    return result


def summarize_raster_image(content):
    with Image.open(io.BytesIO(content)) as image:
        frame_count = getattr(image, "n_frames", 1)
        metadata = [
            f"Image format: {image.format or 'unknown'}",
            f"Dimensions: {image.width} × {image.height} pixels",
            f"Color mode: {image.mode}",
        ]
        if frame_count > 1:
            metadata.append(f"Frames: {frame_count}")
        if image.format == "TIFF":
            tags = [
                f"{tag}: {value}"
                for tag, value in list(image.tag_v2.items())[:12]
                if isinstance(value, (str, int, float))
            ]
            if tags:
                metadata.append("TIFF tags: " + "; ".join(tags))
    metadata.append("Image pixels are not OCR-analyzed.")
    return "\n".join(metadata)


def summarize_svg(content):
    root = ElementTree.fromstring(content)
    text = [
        " ".join(node.itertext()).strip()
        for node in root.iter()
        if node.tag.rsplit("}", 1)[-1] in {"title", "desc", "text"}
    ]
    size = " × ".join(
        str(root.attrib[key])
        for key in ("width", "height")
        if root.attrib.get(key)
    )
    details = [f"SVG dimensions: {size}" if size else "SVG dimensions: not specified"]
    readable_text = "\n".join(value for value in text if value)
    details.append(readable_text or "No text labels found; vector artwork is not OCR-analyzed.")
    return "\n".join(details)


def summarize_zip(content):
    with zipfile.ZipFile(io.BytesIO(content)) as archive:
        entries = [entry for entry in archive.infolist() if not entry.is_dir()]
        if len(entries) > MAX_ARCHIVE_ENTRIES:
            raise ValueError(
                f"Archive contains {len(entries)} files; maximum supported is "
                f"{MAX_ARCHIVE_ENTRIES}."
            )
        total_expanded = sum(entry.file_size for entry in entries)
        if total_expanded > MAX_ARCHIVE_EXPANDED_BYTES:
            raise ValueError("Archive expands beyond the 25 MB processing limit.")

        sections = [f"Files in archive: {len(entries)}"]
        for entry in entries:
            if entry.flag_bits & 0x1:
                sections.append(f"{entry.filename}: skipped (encrypted archive entry)")
                continue
            extension = entry.filename.rsplit(".", 1)[-1].casefold() if "." in entry.filename else ""
            if extension not in UPLOAD_EXTENSIONS or extension == "zip":
                sections.append(f"{entry.filename}: listed; content not extracted")
                continue
            nested_content = archive.read(entry)
            try:
                nested_summary = summarize_upload_content(
                    entry.filename, nested_content, allow_zip=False
                )
                sections.append(
                    f"{entry.filename}:\n{nested_summary[:MAX_REPORT_TEXT_CHARS]}"
                )
            except (
                ValueError, OSError, UnicodeError, PdfReadError,
                InvalidFileException, Image.DecompressionBombError,
                zipfile.BadZipFile,
            ) as exc:
                sections.append(f"{entry.filename}: could not read: {exc}")
        return "\n\n".join(sections)


def summarize_upload_content(filename, content, allow_zip=True):
    extension = filename.rsplit(".", 1)[-1].casefold() if "." in filename else ""
    if len(content) > MAX_UPLOAD_BYTES:
        raise ValueError("File exceeds the 25 MB per-file processing limit.")
    if extension in {"txt", "md", "py"}:
        summary = decode_uploaded_text(content)
        if extension == "py":
            summary = "Python source shown as text only; it is never executed.\n\n" + summary
    elif extension == "html":
        parser = HTMLTextExtractor()
        parser.feed(decode_uploaded_text(content))
        summary = "\n".join(parser.parts) or "No readable text was found in the HTML."
    elif extension == "json":
        try:
            parsed = json.loads(content.decode("utf-8-sig"))
        except (UnicodeDecodeError, json.JSONDecodeError) as exc:
            raise ValueError(f"Invalid JSON: {exc}") from exc
        summary = json.dumps(parsed, indent=2, ensure_ascii=False)
    elif extension == "csv":
        try:
            summary = summarize_csv(content)
        except (pd.errors.ParserError, pd.errors.EmptyDataError, UnicodeDecodeError) as exc:
            raise ValueError(f"Could not read CSV: {exc}") from exc
    elif extension == "xlsx":
        try:
            summary = summarize_workbook(content)
        except (
            OSError, ValueError, KeyError, InvalidFileException, zipfile.BadZipFile
        ) as exc:
            raise ValueError(f"Could not read Excel workbook: {exc}") from exc
    elif extension == "docx":
        try:
            summary = summarize_docx(content)
        except (OSError, KeyError, ElementTree.ParseError, zipfile.BadZipFile) as exc:
            raise ValueError(f"Could not read Word document: {exc}") from exc
    elif extension == "pdf":
        try:
            summary = summarize_pdf(content)
        except (PdfReadError, OSError, ValueError) as exc:
            raise ValueError(f"Could not read PDF: {exc}") from exc
    elif extension == "svg":
        try:
            summary = summarize_svg(content)
        except ElementTree.ParseError as exc:
            raise ValueError(f"Could not read SVG: {exc}") from exc
    elif extension in {"png", "jpg", "jpeg", "tif", "tiff"}:
        try:
            summary = summarize_raster_image(content)
        except (
            OSError, ValueError, UnidentifiedImageError, Image.DecompressionBombError
        ) as exc:
            raise ValueError(f"Could not read image: {exc}") from exc
    elif extension == "zip" and allow_zip:
        try:
            summary = summarize_zip(content)
        except (OSError, zipfile.BadZipFile) as exc:
            raise ValueError(f"Could not read ZIP archive: {exc}") from exc
    else:
        raise ValueError(f"Unsupported file type: .{extension or 'unknown'}")

    return summary[:MAX_REPORT_TEXT_CHARS]


def first_existing(columns, candidates):
    available = {str(column).casefold(): column for column in columns}
    for candidate in candidates:
        column = available.get(str(candidate).casefold())
        if column is not None:
            return column
    return None


def logout():
    for key in ("token", "refresh_token", "user_email"):
        st.session_state.pop(key, None)
    st.rerun()


def load_dashboard_data(token):
    return {
        "health": get_model_health(token),
        "buildings": get_buildings(token),
        "summary": get_portfolio_summary(token),
        "metrics": get_metrics(token),
        "formula": get_formula(token),
    }


def find_value(payload, keys):
    if not isinstance(payload, dict):
        return None

    lowered = {str(key).lower(): value for key, value in payload.items()}
    for key in keys:
        if key.lower() in lowered:
            return lowered[key.lower()]

    for value in payload.values():
        if isinstance(value, dict):
            found = find_value(value, keys)
            if found is not None:
                return found
    return None


def collect_return_period_rows(value):
    rows = []

    if isinstance(value, list):
        for item in value:
            rows.extend(collect_return_period_rows(item))
    elif isinstance(value, dict):
        if "losses" in value and isinstance(value["losses"], dict):
            for period, loss in value["losses"].items():
                rows.append({"return_period": period, "loss_kes": loss})

        keys = value.keys()
        period_key = first_existing(keys, ["return_period", "return_period_years", "rp", "period"])
        loss_key = first_existing(
            keys,
            ["loss_kes", "portfolio_loss_kes", "expected_loss_kes", "loss", "value"],
        )
        if period_key and loss_key:
            rows.append(value)
        else:
            for child in value.values():
                rows.extend(collect_return_period_rows(child))

    return rows


def return_period_df(summary):
    rows = collect_return_period_rows(summary)
    if not rows:
        return pd.DataFrame()

    df = pd.DataFrame(rows)
    period_col = first_existing(df.columns, ["return_period", "return_period_years", "rp", "period"])
    loss_col = first_existing(
        df.columns,
        ["loss_kes", "portfolio_loss_kes", "expected_loss_kes", "loss", "value"],
    )
    if not period_col or not loss_col:
        return pd.DataFrame()

    out = df.copy()
    parsed_periods = out[period_col].astype(str).str.extract(r"(\d+(?:\.\d+)?)", expand=False)
    out["return_period"] = pd.to_numeric(parsed_periods.fillna(out[period_col]), errors="coerce")
    out["loss_kes"] = out[loss_col].map(to_number)
    out = out.dropna(subset=["return_period", "loss_kes"])
    out = out[out["return_period"] > 0]
    return out.sort_values("return_period")


def figure_layout(fig, height=360):
    fig.update_layout(
        height=height,
        margin=dict(l=54, r=20, t=42, b=56),
        paper_bgcolor="white",
        plot_bgcolor="white",
        font=dict(color=TEXT, family="Inter, Arial, sans-serif"),
        legend=dict(orientation="h", yanchor="bottom", y=1.03, xanchor="left", x=0),
        xaxis=dict(gridcolor="#d9e9ee", zerolinecolor="#d9e9ee"),
        yaxis=dict(gridcolor="#d9e9ee", zerolinecolor="#d9e9ee"),
        hoverlabel=dict(
            bgcolor="white",
            bordercolor=LINE,
            font=dict(color=TEXT, family="Inter, Arial, sans-serif"),
        ),
    )
    return fig


def metric_card(icon, label, value, help_text, badge=""):
    badge_html = render_source_badge(badge) if badge else ""
    st.markdown(
        f"""
        <div class="nf-metric">
            <div style="display:flex;justify-content:space-between;gap:12px;">
                <div class="nf-metric-icon">{safe(icon)}</div>
                {badge_html}
            </div>
            <div class="nf-metric-label">{safe(label)}</div>
            <div class="nf-metric-value">{safe(value)}</div>
            <div class="nf-metric-help">{safe(help_text)}</div>
        </div>
        """,
        unsafe_allow_html=True,
    )


def render_source_badge(badge):
    source_badges = {
        "API": ("☁️", "Results from the underwriting service"),
        "API DATA": ("📊", "Portfolio data from the underwriting service"),
    }
    icon, description = source_badges.get(
        badge, (str(badge), str(badge))
    )
    return (
        f'<span class="badge source-icon" role="img" aria-label="{safe(description)}" '
        f'title="{safe(description)}">{safe(icon)}</span>'
    )


def card_start(title=None, badge=None):
    badge_html = render_source_badge(badge) if badge else ""
    title_html = (
        f"<h3 style='margin:0 0 14px;font-size:1.25rem;'>{safe(title)} {badge_html}</h3>"
        if title
        else ""
    )
    st.markdown(f'<div class="nf-card">{title_html}', unsafe_allow_html=True)


def card_end():
    st.markdown("</div>", unsafe_allow_html=True)


def location_column(df):
    return first_existing(
        df.columns,
        [
            "neighbourhood",
            "neighborhood",
            "neighborhood_name",
            "neighbourhood_name",
            "estate",
            "estate_name",
            "area",
            "ward",
            "subcounty",
            "sub_county",
            "sub_county_name",
            "location",
            "place",
            "place_name",
            "address",
        ],
    )


def location_columns(df):
    candidates = [
        "location",
        "neighbourhood",
        "neighborhood",
        "neighborhood_name",
        "neighbourhood_name",
        "area",
        "estate",
        "estate_name",
        "place",
        "place_name",
        "ward",
        "subcounty",
        "sub_county",
        "sub_county_name",
        "address",
    ]
    return list(
        dict.fromkeys(
            column
            for column in (first_existing(df.columns, [name]) for name in candidates)
            if column
        )
    )


def filter_location_records(df, query):
    query = str(query or "").strip()
    columns = location_columns(df)
    if not query or df.empty or not columns:
        return df.copy()

    matches = pd.Series(False, index=df.index)
    for column in columns:
        matches |= df[column].astype("string").str.contains(
            query, case=False, na=False, regex=False
        )
    return df[matches].copy()


def broker_asset_matches(buildings_df):
    available = {str(column).casefold(): column for column in buildings_df.columns}
    identity_candidates = [
        available[name]
        for name in (
            "building_name",
            "property_name",
            "insured_name",
            "client_name",
            "name",
            "address",
            "street_address",
        )
        if name in available
    ]
    loc_cols = location_columns(buildings_df)

    exact_matches = pd.Series(False, index=buildings_df.index)
    for column in identity_candidates:
        values = buildings_df[column].astype("string")
        exact_matches |= values.str.contains(
            r"landmark plaza|landmark realty|lot\s*12.*block\s*2a",
            case=False,
            na=False,
            regex=True,
        )

    location_matches = pd.Series(False, index=buildings_df.index)
    for column in loc_cols:
        location_matches |= buildings_df[column].astype("string").str.contains(
            "Upper Hill", case=False, na=False, regex=False
        )

    if exact_matches.any():
        return buildings_df[exact_matches].copy(), True
    return buildings_df[location_matches].copy(), False


def render_broker_engine_insights(quote):
    if not isinstance(quote, dict):
        st.info("The financial engine returned no quote details.")
        return

    engine_tiv = find_value(
        quote, ["tiv_kes", "insured_value_kes", "total_insured_value_kes"]
    )
    expected_loss = find_value(
        quote,
        [
            "expected_annual_loss_kes",
            "annual_loss_kes",
        ],
    )
    gross_premium = find_value(
        quote, ["gross_premium_kes", "quoted_premium_kes", "premium_kes"]
    )
    deductible = find_value(quote, ["deductible_kes"])
    limit = find_value(quote, ["suggested_limit_kes", "limit_kes"])
    risk_class = find_value(quote, ["risk_class", "risk_category"])
    confidence = find_value(quote, ["confidence", "model_confidence"])

    st.markdown("**Financial-engine indication**")
    metric_values = [
        ("Modelled insured value", engine_tiv, format_kes),
        ("Expected annual loss", expected_loss, format_kes),
        ("Indicated gross premium", gross_premium, format_kes),
        ("Engine deductible", deductible, format_kes),
        ("Suggested limit", limit, format_kes),
        ("Model risk class", risk_class, str),
        ("Model confidence", confidence, str),
    ]
    available_metrics = [
        (label, formatter(value))
        for label, value, formatter in metric_values
        if value is not None
    ]
    for start in range(0, len(available_metrics), 3):
        metric_columns = st.columns(3)
        for column, (label, value) in zip(
            metric_columns, available_metrics[start : start + 3]
        ):
            column.metric(label, value)

    insights = []
    declared_tiv = BROKER_PLACEMENT["insured_value_kes"]
    model_tiv = to_number(engine_tiv)
    if model_tiv is not None and declared_tiv:
        variance = (model_tiv - declared_tiv) / declared_tiv * 100
        if abs(variance) >= 1:
            insights.append(
                f"Validate declared values: the matched model TIV is "
                f"{abs(variance):.1f}% {'above' if variance > 0 else 'below'} "
                "the broker's KES 1.09B figure."
            )
        else:
            insights.append(
                "The matched model TIV is within 1% of the broker's KES 1.09B figure."
            )

    model_limit = to_number(limit)
    if model_limit is not None and declared_tiv:
        limit_share = model_limit / declared_tiv * 100
        if model_limit < declared_tiv:
            insights.append(
                f"The engine's suggested limit covers {limit_share:.1f}% of "
                "the broker's requested full-TIV flood limit."
            )
        else:
            insights.append(
                "The engine's suggested limit meets or exceeds the broker's "
                "requested full-TIV flood limit."
            )

    model_premium = to_number(gross_premium)
    current_premium = BROKER_PLACEMENT["current_premium_kes"]
    if model_premium is not None and current_premium:
        premium_share = model_premium / current_premium * 100
        insights.append(
            f"The engine indication is {premium_share:.1f}% of the broker-reported "
            "current annual premium; this is not a like-for-like comparison because "
            "the current policy excludes flood."
        )

    model_loss = to_number(expected_loss)
    historical_average = (
        BROKER_PLACEMENT["historical_losses_kes"]
        / BROKER_PLACEMENT["historical_loss_years"]
    )
    if model_loss is not None:
        insights.append(
            f"Broker-reported historical average losses are {format_kes(historical_average)} "
            f"per year; the engine's expected annual loss is {format_kes(model_loss)}. "
            "These are different measures, and the broker reports no historical flood losses."
        )

    st.markdown("**Underwriting follow-up**")
    st.markdown(
        "- Flood is excluded from the submitted all-risks cover; explicitly confirm "
        "whether flood cover is being requested and bound.\n"
        "- Confirm the basis of the broker's 5% / KES 5M minimum deductible before "
        "comparing it with the engine deductible.\n"
        "- Review the basement fuel tank, roof membrane, emergency lighting, and "
        "basement wall monitoring noted in the memo."
    )
    if insights:
        st.markdown("**Data-driven insights**")
        for insight in insights:
            st.markdown(f"- {safe(insight)}")
    st.caption(
        "Engine figures are an indication for underwriting review, not a bound quote. "
        "The broker memo is the source for the submitted terms and historical information."
    )


def render_broker_document_review(token):
    st.markdown("#### Analyze a broker report")
    st.caption(
        "Upload the broker's document to the underwriting service. The financial engine "
        "provides calculations; AI explains the returned results for review."
    )
    broker_file = st.file_uploader(
        "Broker report",
        type=["pdf", "docx", "txt", "xlsx", "csv"],
        accept_multiple_files=False,
        key="broker_report_file",
        help="PDF, DOCX, TXT, XLSX, or CSV; maximum 25 MB.",
    )
    portfolio_name = st.text_input(
        "Portfolio",
        value="Commercial Property Portfolio",
        key="broker_report_portfolio_name",
    ).strip()
    if broker_file is None:
        st.caption("Choose a broker report to enable model analysis.")
        return

    content = broker_file.getvalue()
    if len(content) > MAX_UPLOAD_BYTES:
        st.error("This broker report exceeds the 25 MB upload limit.")
        return
    if not portfolio_name:
        st.warning("Enter a portfolio name before submitting the report.")
        return

    fingerprint = sha256(
        broker_file.name.encode("utf-8")
        + b"\0"
        + sha256(content).digest()
        + b"\0"
        + portfolio_name.encode("utf-8")
    ).hexdigest()
    previous = st.session_state.get("broker_document_analysis")
    if previous and previous["fingerprint"] != fingerprint:
        st.session_state.pop("broker_document_analysis", None)
        st.session_state.pop("broker_document_ai_summary", None)
        st.session_state.pop("broker_document_ai_error", None)
        previous = None

    if st.button(
        "Run financial and AI review",
        type="primary",
        disabled=not portfolio_name,
        key="broker_report_analyze",
    ):
        try:
            with st.spinner(
                "Uploading the report and running the available underwriting analysis..."
            ):
                response = upload_document(
                    token,
                    broker_file.name,
                    content,
                    portfolio_name,
                )
            st.session_state["broker_document_analysis"] = {
                "fingerprint": fingerprint,
                "response": response,
            }
            st.session_state.pop("broker_document_ai_summary", None)
            st.session_state.pop("broker_document_ai_error", None)
        except APIError as exc:
            st.session_state["broker_document_analysis"] = {
                "fingerprint": fingerprint,
                "error": str(exc),
                "status_code": exc.status_code,
            }
        previous = st.session_state.get("broker_document_analysis")

    if not previous or previous.get("fingerprint") != fingerprint:
        return
    if previous.get("error"):
        status = (
            f" (HTTP {previous['status_code']})"
            if previous.get("status_code") is not None
            else ""
        )
        st.error(f"Broker report analysis failed{status}: {previous['error']}")
        return

    response = previous.get("response")
    if not isinstance(response, dict):
        st.error(
            "The underwriting service returned an unexpected response; "
            "analysis is unavailable."
        )
        return

    readiness = response.get("model_readiness")
    model_ready = not isinstance(readiness, dict) or readiness.get("ready") is not False
    if model_ready and response.get("model_run") is not None:
        st.success("The backend returned a model run for this report.")
    else:
        st.warning(
            "The report was received, but a completed financial model run was not "
            "returned. The AI can explain extracted information, but no missing "
            "financial losses or return-period values will be invented."
        )

    financial_summary = response.get("financial_summary")
    if not isinstance(financial_summary, dict):
        financial_summary = {}
    model_run = response.get("model_run")
    if not isinstance(model_run, dict):
        model_run = {}

    expected_loss = find_value(
        {"financial_summary": financial_summary, "model_run": model_run},
        ["expected_annual_loss_kes", "portfolio_aal_kes", "aal_kes"],
    )
    aal_rate = find_value(
        {"financial_summary": financial_summary, "model_run": model_run},
        ["aal_rate", "expected_annual_loss_rate"],
    )
    total_tiv = find_value(
        {"financial_summary": financial_summary, "model_run": model_run},
        ["total_insured_value_kes", "tiv_kes", "insured_value_kes"],
    )
    indicative_premium = find_value(
        {"financial_summary": financial_summary, "model_run": model_run},
        ["indicative_gross_premium_kes", "gross_premium_kes"],
    )
    risk_rating = find_value(
        {"financial_summary": financial_summary, "model_run": model_run},
        ["risk_rating", "risk_class", "risk_category"],
    )
    confidence = find_value(
        {"financial_summary": financial_summary, "model_run": model_run},
        ["confidence", "model_confidence"],
    )
    extracted_rows = response.get("building_risk_summary")
    if not isinstance(extracted_rows, list):
        extracted_rows = response.get("extracted_rows", [])
    if not isinstance(extracted_rows, list):
        extracted_rows = []
    risk_rows = []
    for row in extracted_rows:
        if isinstance(row, dict):
            score = to_number(
                find_value(row, ["risk_score", "flood_risk_score"])
            )
            if score is not None:
                risk_rows.append((row, score))
    risk_score = None
    if risk_rows:
        _, risk_score = max(risk_rows, key=lambda item: item[1])

    st.markdown("**Financial-engine results returned by the backend**")
    kpis = [
        ("Total insured value", total_tiv, format_kes),
        ("Expected annual loss", expected_loss, format_kes),
        ("AAL rate", aal_rate, format_pct),
        ("Indicative gross premium", indicative_premium, format_kes),
        ("Risk rating", risk_rating, str),
        ("Highest returned risk score", risk_score, str),
        ("Model confidence", confidence, str),
    ]
    available_kpis = [
        (label, formatter(value))
        for label, value, formatter in kpis
        if value is not None
    ]
    if available_kpis:
        for start in range(0, len(available_kpis), 3):
            columns = st.columns(3)
            for column, (label, value) in zip(
                columns, available_kpis[start : start + 3]
            ):
                column.metric(label, value)
    else:
        st.info("No financial loss, risk, or premium figures were returned.")
    if extracted_rows:
        with st.expander("Extracted building risk records"):
            st.dataframe(
                pd.json_normalize(extracted_rows),
                hide_index=True,
                width="stretch",
            )

    losses_by_return_period = financial_summary.get("losses_per_return_period")
    if not isinstance(losses_by_return_period, dict):
        losses_by_return_period = find_value(
            {
                "model_run": model_run,
                "response": response,
            },
            ["losses_per_return_period", "losses_by_return_period"],
        )
    rp_df = (
        return_period_df({"losses": losses_by_return_period})
        if isinstance(losses_by_return_period, dict)
        else pd.DataFrame()
    )
    if not rp_df.empty:
        rp_df = rp_df.drop_duplicates(
            subset=["return_period"], keep="last"
        ).sort_values("return_period")
        chart = go.Figure(
            go.Scatter(
                x=rp_df["return_period"],
                y=rp_df["loss_kes"],
                mode="lines+markers",
                line={"color": TEAL, "width": 3, "shape": "spline"},
                marker={"size": 8, "color": AQUA},
                customdata=(100 / rp_df["return_period"]).tolist(),
                hovertemplate=(
                    "Return period: %{x:,.0f} years"
                    "<br>Annual exceedance probability: %{customdata:.2f}%"
                    "<br>Loss: KES %{y:,.0f}<extra></extra>"
                ),
            )
        )
        chart.update_layout(
            height=300,
            margin={"l": 10, "r": 10, "t": 20, "b": 10},
            xaxis_title="Return period (years)",
            yaxis_title="Loss (KES)",
            hovermode="x unified",
        )
        st.markdown("**Loss by return period**")
        st.plotly_chart(
            chart,
            width="stretch",
            key="broker_document_return_period_chart",
        )
    else:
        st.info(
            "No loss-return-period curve can be plotted because this upload "
            "response has no modelled loss values."
        )

    hazard_fields = [
        ("Common hazard", "common_hazard"),
        ("Occasional hazard", "occasional_hazard"),
        ("Moderate hazard", "moderate_hazard"),
        ("Severe hazard", "severe_hazard"),
        ("Extreme hazard", "extreme_hazard"),
    ]
    hazard_chart_rows = []
    for row in extracted_rows:
        if not isinstance(row, dict):
            continue
        row_name = (
            row.get("building_id")
            or row.get("building_name")
            or row.get("neighbourhood")
            or "Extracted exposure"
        )
        for label, field in hazard_fields:
            value = to_number(row.get(field))
            if value is not None:
                hazard_chart_rows.append(
                    {
                        "Building / exposure": str(row_name),
                        "Hazard input": label,
                        "Extracted value": value,
                    }
                )
    if hazard_chart_rows:
        st.markdown("**Extracted hazard inputs (not modelled losses)**")
        st.caption(
            "These are source/extraction input values, not financial-loss estimates "
            "or calibrated probabilities."
        )
        hazard_chart = px.bar(
            pd.DataFrame(hazard_chart_rows),
            x="Hazard input",
            y="Extracted value",
            color="Building / exposure",
            barmode="group",
            color_discrete_sequence=CLASS_COLORS,
            hover_data={"Extracted value": ":.3f"},
        )
        hazard_chart.update_layout(
            height=300,
            margin={"l": 10, "r": 10, "t": 20, "b": 10},
            yaxis_title="Raw extracted factor value",
            xaxis_title="Hazard scenario",
            legend_title_text="Building / exposure",
        )
        st.plotly_chart(
            hazard_chart,
            width="stretch",
            key="broker_document_hazard_inputs_chart",
        )

    warnings = response.get("warnings")
    limitations = response.get("limitations")
    assumptions = response.get("assumptions")
    extracted_data = response.get("extracted_data")
    if not isinstance(extracted_data, dict):
        extracted_data = {}
    if assumptions is None:
        assumptions = financial_summary.get("assumptions")
    if limitations is None:
        limitations = financial_summary.get("limitations")

    with st.expander("Model readiness, assumptions & limitations", expanded=True):
        st.markdown("**Readiness and source warnings**")
        if isinstance(readiness, dict):
            st.markdown(
                f"Model ready: **{'Yes' if readiness.get('ready') else 'No'}**"
            )
        if warnings:
            warning_items = warnings if isinstance(warnings, list) else [warnings]
            for warning in warning_items:
                st.markdown(f"- {safe(warning)}")
        if isinstance(readiness, dict) and readiness.get("issues"):
            st.markdown("**Items blocking or limiting the model run**")
            for issue in readiness["issues"]:
                st.markdown(f"- {safe(issue)}")
                issue_text = str(issue).casefold()
                if "ai_confidence" in issue_text and (
                    "valid number" in issue_text or "float" in issue_text
                ):
                    st.warning(
                        "Backend fix required: the extraction pipeline is sending "
                        "a text confidence label where the model schema expects a "
                        "number. Update the extractor/schema mapping and rerun the "
                        "model; do not substitute an arbitrary confidence value."
                    )
        if assumptions:
            st.markdown("**Assumptions returned by the backend**")
            st.write(assumptions)
        else:
            extraction_notes = [
                row.get("extraction_notes")
                for row in extracted_rows
                if isinstance(row, dict) and row.get("extraction_notes")
            ]
            st.markdown("**Assumptions / provenance**")
            st.write(
                "No formal model assumptions were returned. Treat document "
                "extractions and AI-derived fields as unverified until checked "
                "against the source report."
            )
            for note in extraction_notes:
                st.markdown(f"- {safe(note)}")
        if limitations:
            st.markdown("**Limitations returned by the backend**")
            st.write(limitations)
        else:
            st.markdown("**Current analysis limitations**")
            limitation_items = []
            if not model_run:
                limitation_items.append(
                    "No completed model run was returned; expected loss, premium, "
                    "and return-period losses are not calculated."
                )
            if rp_df.empty:
                limitation_items.append(
                    "No loss-return-period data was returned, so an EP/loss curve "
                    "cannot be calculated from this response."
                )
            limitation_items.append(
                "Extracted values are subject to source-document quality and "
                "require underwriting verification."
            )
            for limitation in limitation_items:
                st.markdown(f"- {limitation}")

    recommendations = extracted_data.get("underwriting_recommendations")
    if isinstance(recommendations, dict):
        with st.expander("Extracted underwriting recommendations"):
            for label in (
                "overall_risk_profile",
                "final_recommendation",
            ):
                value = recommendations.get(label)
                if value:
                    st.markdown(f"**{label.replace('_', ' ').title()}:** {safe(value)}")
            for label in (
                "positive_factors",
                "areas_requiring_attention",
                "required_conditions",
            ):
                values = recommendations.get(label)
                if isinstance(values, list) and values:
                    st.markdown(f"**{label.replace('_', ' ').title()}**")
                    for value in values:
                        st.markdown(f"- {safe(value)}")

    ai_portfolio_summary = {
        **financial_summary,
        "portfolio_name": portfolio_name,
        "document_type": response.get("document_type"),
        "model_readiness": readiness,
        "model_run": model_run or None,
        "warnings": warnings,
        "assumptions": assumptions,
        "limitations": limitations,
        "document_summary": response.get("summary"),
    }
    if st.button(
        "Generate AI underwriter briefing",
        type="primary",
        key="broker_report_ai_summary",
    ):
        try:
            with st.spinner("Preparing an underwriter-friendly briefing..."):
                ai_response = request_uploaded_document_summary(
                    token,
                    ai_portfolio_summary,
                    extracted_rows,
                )
            st.session_state["broker_document_ai_summary"] = {
                "fingerprint": fingerprint,
                "response": ai_response,
            }
            st.session_state.pop("broker_document_ai_error", None)
        except APIError as exc:
            st.session_state["broker_document_ai_error"] = {
                "fingerprint": fingerprint,
                "message": str(exc),
                "status_code": exc.status_code,
            }
            st.session_state.pop("broker_document_ai_summary", None)

    ai_error = st.session_state.get("broker_document_ai_error")
    ai_summary_text = None
    if ai_error and ai_error.get("fingerprint") == fingerprint:
        status = (
            f" (HTTP {ai_error['status_code']})"
            if ai_error.get("status_code") is not None
            else ""
        )
        st.error(f"AI briefing failed{status}: {ai_error['message']}")
        if (
            ai_error.get("status_code") == 502
            or "interpretation service is unavailable"
            in ai_error["message"].casefold()
        ):
            st.info(
                "The report upload succeeded; the AI provider behind the backend "
                "proxy is unavailable. Ask the backend owner to start the "
                "interpretation service (port 5001) and verify the proxy's upstream "
                "URL, then select **Generate AI underwriter briefing** to retry. "
                "The report data and calculations below remain available."
            )
    ai_result = st.session_state.get("broker_document_ai_summary")
    if ai_result and ai_result.get("fingerprint") == fingerprint:
        ai_response = ai_result.get("response")
        ai_summary_text = (
            ai_response.get("summary")
            if isinstance(ai_response, dict)
            else ai_response
        )
        if isinstance(ai_summary_text, str) and ai_summary_text.strip():
            st.markdown("**AI explanation for the underwriter**")
            st.markdown(ai_summary_text)
            if isinstance(ai_response, dict) and ai_response.get("source"):
                st.caption(f"AI source: {safe(ai_response['source'])}")
        else:
            st.info("The AI service returned no summary text.")
            st.json(ai_response, expanded=True)
        if not model_ready or not response.get("model_run"):
            st.warning(
                "This AI explanation is based on the extracted report only; a "
                "completed financial model run is still required for calculated "
                "expected loss and return-period curves."
            )

    if ai_error and ai_error.get("fingerprint") == fingerprint:
        st.markdown("**Available source-backed report summary (not AI-generated)**")
        document_summary = response.get("summary")
        if document_summary:
            st.markdown(f"- **Document:** {safe(document_summary)}")
        st.markdown(f"- **Portfolio:** {safe(portfolio_name)}")
        st.markdown(f"- **Total insured value:** {format_kes(total_tiv)}")
        st.markdown(f"- **Expected annual loss:** {format_kes(expected_loss)}")
        st.markdown(f"- **Indicative gross premium:** {format_kes(indicative_premium)}")
        if isinstance(readiness, dict):
            st.markdown(
                f"- **Model readiness:** "
                f"{'Ready' if readiness.get('ready') is True else 'Not ready'}"
            )
            for issue in readiness.get("issues", []):
                st.markdown(f"  - **Model issue:** {safe(issue)}")
        for warning in warnings if isinstance(warnings, list) else [warnings]:
            if warning:
                st.markdown(f"- **Warning:** {safe(warning)}")

    report_sections = [
        "# Broker report underwriting analysis",
        "",
        f"- Portfolio: {portfolio_name}",
        f"- Source file: {broker_file.name}",
        f"- Document type: {response.get('document_type') or 'Not returned'}",
        f"- Total insured value: {format_kes(total_tiv)}",
        f"- Expected annual loss: {format_kes(expected_loss)}",
        f"- AAL rate: {format_pct(aal_rate)}",
        f"- Indicative gross premium: {format_kes(indicative_premium)}",
        f"- Risk rating: {risk_rating if risk_rating is not None else 'Not returned'}",
        f"- Model confidence: {confidence if confidence is not None else 'Not returned'}",
        "",
        "## Model readiness",
        "",
        json.dumps(readiness, indent=2, ensure_ascii=False, default=str)
        if readiness is not None
        else "No readiness information was returned.",
        "",
        "## Loss by return period",
        "",
        rp_df[["return_period", "loss_kes"]].to_json(
            orient="records", indent=2
        )
        if not rp_df.empty
        else "No modelled return-period losses were returned; no loss curve is available.",
        "",
        "## Extracted hazard inputs (not modelled losses)",
        "",
        json.dumps(hazard_chart_rows, indent=2, ensure_ascii=False)
        if hazard_chart_rows
        else "No recognized hazard-input fields were returned.",
        "",
        "## Warnings",
        "",
        json.dumps(warnings, indent=2, ensure_ascii=False, default=str)
        if warnings
        else "No warnings were returned.",
        "",
        "## Assumptions and limitations",
        "",
        json.dumps(
            {
                "assumptions": assumptions,
                "limitations": limitations,
                "extraction_notes": [
                    row.get("extraction_notes")
                    for row in extracted_rows
                    if isinstance(row, dict) and row.get("extraction_notes")
                ],
            },
            indent=2,
            ensure_ascii=False,
            default=str,
        ),
        "",
        "## Extracted underwriting recommendations",
        "",
        json.dumps(recommendations, indent=2, ensure_ascii=False, default=str)
        if recommendations
        else "No recommendations were returned.",
        "",
        "## AI explanation",
        "",
        ai_summary_text
        if isinstance(ai_summary_text, str) and ai_summary_text.strip()
        else "AI interpretation was unavailable; this report contains backend-returned data only.",
        "",
        "This report is for underwriting review. Extracted fields require validation "
        "against the source document. No missing model outputs have been estimated.",
    ]
    st.download_button(
        "Download underwriting analysis (.md)",
        "\n".join(report_sections),
        file_name="broker-underwriting-analysis.md",
        mime="text/markdown",
        width="stretch",
        key="broker_document_analysis_download",
    )


def mark_broker_notification_read():
    st.session_state["broker_notification_read"] = True


def render_broker_notification(token, buildings_df):
    is_unread = not st.session_state.get("broker_notification_read", False)
    if is_unread:
        st.markdown(
            """
            <style>
            div.st-key-broker_notification_popover
            [data-testid="stPopover"] > button::after {
                position: absolute;
                top: 5px;
                right: 5px;
                width: 11px;
                height: 11px;
                border: 2px solid #fff;
                border-radius: 50%;
                background: #ef4444;
                content: "";
            }
            </style>
            """,
            unsafe_allow_html=True,
        )

    with st.popover(
        "Notifications",
        icon=":material/notifications:",
        help="Open broker and cedant messages",
        key="broker_notification_popover",
        width=640,
    ):
        st.markdown("#### Broker placement")
        st.caption(
            f"{BROKER_PLACEMENT['reference']} · Issued {BROKER_PLACEMENT['issued']}"
        )
        st.markdown(
            f"""
            <div class="broker-message-preview">
                <p><b>{safe(BROKER_PLACEMENT['client'])}</b><br>
                {safe(BROKER_PLACEMENT['location'])}</p>
            </div>
            """,
            unsafe_allow_html=True,
        )
        st.warning(
            f"Review by {BROKER_PLACEMENT['expiry']}: flood is excluded from the "
            "submitted all-risks cover; optional flood terms are proposed."
        )
        st.markdown(
            f"**Broker risk:** {BROKER_PLACEMENT['broker_risk_profile']} · "
            f"**Declared TIV:** {format_kes(BROKER_PLACEMENT['insured_value_kes'])}"
        )

        with st.expander("Broker memo and financial-engine review"):
            st.markdown(
                f"- **Placement reference:** {BROKER_PLACEMENT['reference']}\n"
                f"- **Coverage:** {BROKER_PLACEMENT['coverage']}\n"
                f"- **Flood limit requested:** {BROKER_PLACEMENT['flood_limit']}\n"
                f"- **Flood deductible stated:** {BROKER_PLACEMENT['flood_deductible']}\n"
                f"- **Current annual premium:** {format_kes(BROKER_PLACEMENT['current_premium_kes'])}\n"
                f"- **Reported historical losses:** {format_kes(BROKER_PLACEMENT['historical_losses_kes'])} "
                f"over {BROKER_PLACEMENT['historical_loss_years']} years; no flood losses "
                "reported since 2015."
            )
            st.markdown("**Items for underwriter review**")
            for item in BROKER_PLACEMENT["attention_items"]:
                st.markdown(f"- {safe(item)}")

            candidates, exact_match = broker_asset_matches(buildings_df)
            building_col = first_existing(
                candidates.columns, ["building_id", "id"]
            )
            location_col = location_column(candidates)
            name_col = first_existing(
                candidates.columns,
                ["building_name", "property_name", "name", "address"],
            )

            if candidates.empty:
                st.info(
                    "No model building matched Landmark Plaza or Upper Hill. "
                    "Add the property to the available portfolio records before requesting a "
                    "financial-engine indication; no quote has been inferred from "
                    "the broker memo alone."
                )
            elif not building_col:
                st.info(
                    "The matching property records have no building ID, so the existing "
                    "financial-engine quote endpoint cannot be called."
                )
            else:
                if exact_match:
                    st.success(
                        "A model asset matched the property name or street address."
                    )
                else:
                    st.warning(
                        "These property records match Upper Hill by location only. "
                        "Confirm the exact insured asset before running the engine."
                    )

                candidates = candidates.reset_index(drop=True)

                def candidate_label(index):
                    row = candidates.iloc[index]
                    building_id = row[building_col]
                    name = (
                        row[name_col]
                        if name_col and pd.notna(row[name_col])
                        else row[location_col]
                        if location_col and pd.notna(row[location_col])
                        else "Model building"
                    )
                    return f"{building_id} · {name}"

                selected_index = st.selectbox(
                    "Match broker memo to model asset",
                    range(len(candidates)),
                    format_func=candidate_label,
                    key="broker_review_building",
                )
                selected_row = candidates.iloc[selected_index]
                selected_id = str(selected_row[building_col])
                confirmed = exact_match or st.checkbox(
                    "I confirm this is the Landmark Plaza insured property.",
                    key="broker_review_asset_confirmed",
                )

                if st.button(
                    "Generate financial-engine insight",
                    type="primary",
                    disabled=not confirmed,
                    key="broker_review_run",
                ):
                    result = {"building_id": selected_id}
                    with st.spinner("Requesting the model quote..."):
                        try:
                            result["quote"] = get_building_quote(token, selected_id)
                        except APIError as exc:
                            if exc.status_code != 404:
                                result["error"] = str(exc)
                            else:
                                payload = quote_payload_from_row(selected_row)
                                required = {
                                    "tiv_kes",
                                    "expected_annual_loss_kes",
                                    "risk_class",
                                    "confidence",
                                }
                                if not required.issubset(payload):
                                    result["error"] = (
                                        "The model asset has no stored quote and is "
                                        "missing fields required by the financial "
                                        "engine: "
                                        + ", ".join(sorted(required - payload.keys()))
                                    )
                                else:
                                    try:
                                        result["quote"] = create_quote(token, payload)
                                    except APIError as quote_exc:
                                        result["error"] = str(quote_exc)
                    st.session_state["broker_financial_review"] = result

                result = st.session_state.get("broker_financial_review", {})
                if result.get("building_id") == selected_id:
                    if result.get("error"):
                        st.error(result["error"])
                    elif "quote" in result:
                        if result["quote"] is None:
                            st.warning(
                                "The financial engine returned an empty response; "
                                "no insight was generated."
                            )
                        else:
                            render_broker_engine_insights(result["quote"])

        render_broker_document_review(token)

        st.caption(
            "Preview message · Broker and cedant messages are not connected to a live inbox yet."
        )
        if is_unread:
            st.button(
                "Mark as read",
                key="broker_notification_mark_read",
                width="stretch",
                on_click=mark_broker_notification_read,
            )


def set_dashboard_page(page_name):
    st.session_state["dashboard_navigation"] = page_name


def interpretation_model_output(summary, metrics, formula, buildings_df):
    building_records = json.loads(buildings_df.to_json(orient="records"))
    return {
        "portfolio_summary": summary,
        "metrics": metrics,
        "formula": formula,
        "buildings": building_records,
    }


def render_interpretation_response(response):
    if isinstance(response, str):
        st.markdown(response)
    elif response is None:
        st.info("The interpretation service returned an empty response.")
    else:
        st.json(response, expanded=True)


def interpretation_response_text(response, keys):
    if isinstance(response, str):
        return response
    if isinstance(response, dict):
        for key in keys:
            value = response.get(key)
            if isinstance(value, str) and value.strip():
                return value
    return None


def render_interpretation_summary(summary, metrics, formula, buildings_df):
    card_start("Underwriter interpretation", "AI ASSISTED")
    st.caption(
        "Generate an explanation from the current model response. The data is sent "
        "to the separate interpretation service; it is not a quote or underwriting decision."
    )
    if st.button(
        "Generate dashboard summary",
        type="primary",
        key="generate_interpretation_summary",
    ):
        model_output = interpretation_model_output(
            summary, metrics, formula, buildings_df
        )
        try:
            with st.spinner("Generating an underwriter-friendly summary..."):
                response = request_interpretation_summary(model_output)
            st.session_state["interpretation_summary_result"] = response
            st.session_state.pop("interpretation_summary_error", None)
        except APIError as exc:
            st.session_state["interpretation_summary_error"] = str(exc)
            st.session_state.pop("interpretation_summary_result", None)

    if st.session_state.get("interpretation_summary_error"):
        st.error(st.session_state["interpretation_summary_error"])
    elif "interpretation_summary_result" in st.session_state:
        render_interpretation_response(
            st.session_state["interpretation_summary_result"]
        )
    card_end()


def render_interpretation_page(summary, metrics, formula, buildings_df):
    st.markdown("<h1>Interpretation & questions</h1>", unsafe_allow_html=True)
    st.markdown(
        "<div class='hero-subtitle'>Get an underwriter-friendly summary and ask "
        "questions grounded in the latest model output.</div>",
        unsafe_allow_html=True,
    )
    st.info(
        "Interpretations are generated by a separate service from backend model "
        "output. Treat responses as explanations, not a substitute for underwriting review."
    )
    st.caption(f"Interpretation service: {safe(get_interpretation_base_url())}")

    if st.button("Check interpretation service", key="interpretation_health_check"):
        try:
            with st.spinner("Checking interpretation service..."):
                health = check_interpretation_health()
            st.session_state["interpretation_health_result"] = health
            st.session_state.pop("interpretation_health_error", None)
        except APIError as exc:
            st.session_state["interpretation_health_error"] = str(exc)
            st.session_state.pop("interpretation_health_result", None)

    if st.session_state.get("interpretation_health_error"):
        st.error(st.session_state["interpretation_health_error"])
    elif "interpretation_health_result" in st.session_state:
        st.success("Interpretation service responded.")
        render_interpretation_response(
            st.session_state["interpretation_health_result"]
        )

    model_output = interpretation_model_output(
        summary, metrics, formula, buildings_df
    )
    with st.expander("Model output sent for interpretation"):
        st.json(model_output, expanded=False)

    summary_tab, question_tab = st.tabs(
        ["Dashboard summary", "Ask a question"]
    )
    with summary_tab:
        render_interpretation_summary(
            summary, metrics, formula, buildings_df
        )
    with question_tab:
        question = st.text_area(
            "Question",
            placeholder="For example: Which return period has the largest loss, and what does that mean for the portfolio?",
            max_chars=1000,
            key="interpretation_question",
        )
        if st.button(
            "Ask AI about the model",
            type="primary",
            disabled=not question.strip(),
            key="ask_interpretation_question",
        ):
            try:
                with st.spinner("Preparing an answer from the model output..."):
                    response = ask_interpretation_question(
                        model_output, question.strip()
                    )
                st.session_state["interpretation_question_result"] = {
                    "question": question.strip(),
                    "response": response,
                }
                st.session_state.pop("interpretation_question_error", None)
            except APIError as exc:
                st.session_state["interpretation_question_error"] = str(exc)
                st.session_state.pop("interpretation_question_result", None)

        if st.session_state.get("interpretation_question_error"):
            st.error(st.session_state["interpretation_question_error"])
        else:
            result = st.session_state.get("interpretation_question_result")
            if result:
                st.markdown(f"**Question:** {safe(result['question'])}")
                render_interpretation_response(result["response"])


def render_dashboard_header(token, buildings_df):
    header_col, report_col, notification_col = st.columns(
        [7, 2, 1],
        vertical_alignment="center",
    )
    with header_col:
        st.markdown(
            f"""
            <div class="welcome-date">{datetime.now().strftime("%A, %d %B %Y")}</div>
            <h1 class="welcome-greeting">Welcome back, Underwriter</h1>
            <div class="welcome-subtitle">
                Here is your Nairobi flood portfolio at a glance.
            </div>
            """,
            unsafe_allow_html=True,
        )
    with report_col:
        st.button(
            "Generate report",
            key="header_generate_report",
            width="stretch",
            on_click=set_dashboard_page,
            args=("Generate Report",),
            help="Open the report generation page",
        )
    with notification_col:
        render_broker_notification(token, buildings_df)


def render_sidebar():
    with st.sidebar:
        st.markdown(
            """
            <div style="padding:8px 8px 24px;">
                <div style="display:flex;gap:10px;align-items:center;">
                    <div style="color:#16b9b5;font-size:28px;font-weight:800;">~~~</div>
                    <div>
                        <div style="font-size:1.35rem;font-weight:800;color:white;">Nairobi Flood</div>
                        <div style="font-size:.88rem;color:rgba(255,255,255,.65);">Urban pluvial risk - Team A</div>
                    </div>
                </div>
            </div>
            """,
            unsafe_allow_html=True,
        )
        page = st.radio(
            "Navigation",
            [
                "Overview",
                "EP Curve",
                "Hotspots",
                "Drainage Reports",
                "Generate Report",
                "Find Quote",
                "Methodology",
                "BULK UPLOAD",
                "Data Explorer",
                "Interpretation",
            ],
            format_func=lambda label: {
                "Overview": "🏠  Overview",
                "EP Curve": "📈  Loss by return period",
                "Hotspots": "📍  Flood hotspots",
                "Drainage Reports": "🌧️  Drainage reports",
                "Generate Report": "📄  Generate report",
                "Find Quote": "🔎  Find a quote",
                "Methodology": "🧭  Methodology",
                "BULK UPLOAD": "📤  Bulk upload",
                "Data Explorer": "🔌  Data explorer",
                "Interpretation": "✨  AI interpretation",
            }.get(label, label),
            label_visibility="collapsed",
            key="dashboard_navigation",
        )
        st.markdown("<div style='height:220px'></div>", unsafe_allow_html=True)
        st.caption(st.session_state.get("user_email", "Signed in"))
        if st.button("Log out", width="stretch"):
            logout()
    return page


def render_overview(summary, metrics, formula, buildings_df, rp_df):
    st.markdown("<h2>Portfolio overview</h2>", unsafe_allow_html=True)

    tiv_col = first_existing(buildings_df.columns, ["tiv_kes", "insured_value_kes", "value_kes"])
    loss_col = first_existing(
        buildings_df.columns,
        ["expected_annual_loss_kes", "annual_loss_kes", "loss_kes"],
    )
    risk_col = first_existing(buildings_df.columns, ["risk_class", "risk_category"])

    total_tiv = find_value(summary, ["total_tiv_kes", "tiv_kes", "insured_value_kes", "total_insured_value_kes"])
    if total_tiv is None and tiv_col:
        total_tiv = buildings_df[tiv_col].map(to_number).sum()

    building_count = find_value(summary, ["building_count", "buildings_count", "total_buildings", "count"])
    if building_count is None:
        building_count = len(buildings_df)

    loss_100 = None
    if not rp_df.empty:
        matching = rp_df[rp_df["return_period"] == 100]
        if not matching.empty:
            loss_100 = matching.iloc[0]["loss_kes"]
    if loss_100 is None:
        loss_100 = find_value(summary, ["loss_100_kes", "loss_100yr_kes", "pml_100_kes"])

    loss_rate = to_number(loss_100) / to_number(total_tiv) if to_number(loss_100) and to_number(total_tiv) else None
    high_risk = 0
    if risk_col:
        high_risk = buildings_df[risk_col].astype(str).str.lower().str.contains("high|severe|extreme").sum()
    elif loss_col and not buildings_df.empty:
        high_risk = int((buildings_df[loss_col].map(to_number) > buildings_df[loss_col].map(to_number).median()).sum())

    c1, c2, c3, c4 = st.columns(4)
    with c1:
        metric_card("↗", "Loss in a 100-year event", format_kes(loss_100), "About a 1% chance in any year", "API")
    with c2:
        metric_card("▥", "Insured value", format_kes(total_tiv), f"{format_number(building_count)} buildings", "API")
    with c3:
        metric_card("⌖", "Higher-risk assets", f"{format_number(high_risk)} of {format_number(building_count)}", "From building results", "API")
    with c4:
        metric_card("↗", "100-yr loss as % of value", format_pct(loss_rate), "Share of portfolio lost", "API")

    left, right = st.columns([2.1, 1])
    with left:
        card_start("The story in three lines")
        if not rp_df.empty and len(rp_df) > 1:
            growth = rp_df.iloc[-1]["loss_kes"] / rp_df.iloc[0]["loss_kes"] if rp_df.iloc[0]["loss_kes"] else None
            first_line = f"Losses grow fast with rarer storms. The largest return-period loss is {growth:.1f}x the smallest shown." if growth else "Losses grow fast with rarer storms."
        else:
            first_line = "Losses grow fast with rarer storms when return-period results are available."

        housing_col = first_existing(buildings_df.columns, ["housing_class", "class", "building_class"])
        if housing_col and loss_col and not buildings_df.empty:
            class_loss = (
                buildings_df.assign(_loss=buildings_df[loss_col].map(to_number))
                .groupby(housing_col)["_loss"]
                .sum()
                .sort_values(ascending=False)
            )
            second_line = f"{class_loss.index[0]} carries the largest modeled loss in the current portfolio."
        else:
            second_line = "Building-level loss and class fields will drive the portfolio class story when available."

        st.markdown(
            f"""
            <div class="story-card">
                <p><b>1. {safe(first_line)}</b></p>
                <p><b>2. {safe(second_line)}</b></p>
                <p><b>3. Use the quote search by location.</b> Search places such as Kasarani or Lavington to find matching assets and request quotes from the backend.</p>
            </div>
            """,
            unsafe_allow_html=True,
        )
        card_end()

    with right:
        card_start("Explore")
        explore_pages = [
            ("Loss by return period   →", "EP Curve", "explore_ep_curve"),
            ("Hotspot map   →", "Hotspots", "explore_hotspots"),
            ("Find quote by location   →", "Find Quote", "explore_find_quote"),
            ("Methodology & limitations   →", "Methodology", "explore_methodology"),
        ]
        for label, destination, key in explore_pages:
            st.button(
                label,
                key=key,
                width="stretch",
                on_click=set_dashboard_page,
                args=(destination,),
            )
        card_end()

    if metrics:
        st.markdown(
            "<div class='notice'><b>✅ Portfolio data loaded.</b> "
            "Figures reflect the latest results available from the underwriting service.</div>",
            unsafe_allow_html=True,
        )

    render_location_chart(buildings_df)
    render_interpretation_summary(
        summary, metrics, formula, buildings_df
    )


def render_location_chart(buildings_df):
    card_start("Portfolio by neighbourhood", "API DATA")
    loc_col = location_column(buildings_df)
    value_col = first_existing(
        buildings_df.columns, ["tiv_kes", "insured_value_kes", "value_kes"]
    )
    loss_col = first_existing(
        buildings_df.columns,
        ["expected_annual_loss_kes", "annual_loss_kes", "loss_kes"],
    )
    loss_label = (
        "Reported loss"
        if loss_col and loss_col.casefold() == "loss_kes"
        else "Expected annual loss"
    )

    measures = {"Building count": None}
    if value_col:
        measures["Insured value"] = value_col
    if loss_col:
        measures[loss_label] = loss_col

    if buildings_df.empty or not loc_col:
        st.info(
            "The portfolio data needs a neighbourhood or location field to chart this breakdown."
        )
    else:
        measure = st.selectbox(
            "Compare neighbourhoods by",
            list(measures),
            key="overview_location_measure",
        )
        chart_df = buildings_df[[loc_col]].copy()
        chart_df["_location"] = chart_df[loc_col].fillna("Unspecified").astype(str)

        measure_col = measures[measure]
        if measure_col is None:
            chart_df["_value"] = 1
            value_label = "Buildings"
            aggregation = "count"
        else:
            chart_df["_value"] = buildings_df[measure_col].map(to_number)
            chart_df = chart_df.dropna(subset=["_value"])
            value_label = measure
            aggregation = "sum"

        grouped = (
            chart_df.groupby("_location")["_value"]
            .agg(aggregation)
            .sort_values(ascending=False)
            .head(15)
            .sort_values()
            .reset_index()
        )
        if grouped.empty:
            st.info(f"No usable {measure.lower()} values were returned for the neighbourhood chart.")
        else:
            fig = px.bar(
                grouped,
                x="_value",
                y="_location",
                orientation="h",
                color="_value",
                color_continuous_scale=["#b9e6e3", TEAL],
                labels={"_value": value_label, "_location": "Neighbourhood"},
                hover_data={"_value": ":,.0f"},
            )
            fig.update_layout(coloraxis_showscale=False, showlegend=False)
            fig.update_yaxes(title=None)
            st.plotly_chart(figure_layout(fig, height=390), width="stretch")
            st.caption(
                "Choose a measure to regroup the latest available property records."
            )
    card_end()


def render_ep_curve(summary, buildings_df, rp_df):
    st.markdown("<h1>Loss by return period</h1>", unsafe_allow_html=True)
    st.markdown("<div class='hero-subtitle'>How much the portfolio could lose in floods of different rarity.</div>", unsafe_allow_html=True)
    st.markdown(
        "<div class='badge-row'><span class='badge synthetic' "
        "title='Portfolio data from the underwriting service' "
        "aria-label='Portfolio data from the underwriting service'>📊</span></div>",
        unsafe_allow_html=True,
    )

    top_left, top_right = st.columns([2.2, 1])
    with top_left:
        card_start("Loss vs return period", "API")
        if rp_df.empty:
            st.info("The portfolio summary did not include return-period loss rows.")
        else:
            curve_df = (
                rp_df[["return_period", "loss_kes"]]
                .dropna()
                .sort_values("return_period")
                .reset_index(drop=True)
            )
            curve_df = curve_df[
                (curve_df["return_period"] > 0) & (curve_df["loss_kes"] >= 0)
            ]
            curve_df = curve_df.drop_duplicates(
                subset=["return_period"], keep="last"
            )
            curve_df = curve_df.sort_values("return_period").reset_index(drop=True)
            maximum_loss = float(curve_df["loss_kes"].max())
            if maximum_loss >= 1_000_000:
                axis_unit = 1_000_000
                axis_suffix = "M"
            elif maximum_loss >= 1_000:
                axis_unit = 1_000
                axis_suffix = "K"
            else:
                axis_unit = 1
                axis_suffix = ""
            raw_step = maximum_loss / axis_unit / 4 if maximum_loss else 1
            magnitude = 10 ** int(math.floor(math.log10(raw_step))) if raw_step else 1
            normalized_step = raw_step / magnitude
            nice_step = next(
                value for value in (1, 2, 2.5, 3, 5, 10)
                if value >= normalized_step
            ) * magnitude
            step = nice_step * axis_unit
            axis_maximum = math.ceil(maximum_loss / step) * step if maximum_loss else step
            tick_count = int(round(axis_maximum / step))
            loss_ticks = [index * step for index in range(tick_count + 1)]
            loss_tick_labels = [
                f"{value / axis_unit:,.0f}{axis_suffix}"
                for value in loss_ticks
            ]
            fig = go.Figure()
            fig.add_trace(
                go.Scatter(
                    x=curve_df["return_period"],
                    y=curve_df["loss_kes"],
                    mode="lines+markers",
                    line=dict(color=TEAL, width=4, shape="spline", smoothing=0.35),
                    marker=dict(size=11, color=AQUA, line=dict(color=TEAL, width=3)),
                    customdata=100 / curve_df["return_period"],
                    hovertemplate=(
                        "Number of years: %{x:g}"
                        "<br>Annual chance: %{customdata:.2f}%"
                        "<br>Loss: KES %{y:,.0f}<extra></extra>"
                    ),
                )
            )
            fig.update_layout(showlegend=False, hovermode="closest")
            fig.update_xaxes(
                title="Return period (years)",
                tickmode="array",
                tickvals=curve_df["return_period"].tolist(),
                ticktext=[f"{period:g}" for period in curve_df["return_period"]],
                rangemode="tozero",
            )
            fig.update_yaxes(
                title="Loss (KES)",
                tickmode="array",
                tickvals=loss_ticks,
                ticktext=loss_tick_labels,
                rangemode="tozero",
            )
            st.plotly_chart(figure_layout(fig, height=390), width="stretch")
            st.caption(
                "Hover a point to see the number of years and loss amount from the data. "
                "Each point is a backend-returned portfolio loss estimate."
            )
        card_end()

    with top_right:
        card_start("Loss table")
        if rp_df.empty:
            st.info("No return-period table data returned.")
        else:
            table = rp_df[["return_period", "loss_kes"]].copy().sort_values("return_period")
            table["Return period"] = table["return_period"].map(lambda x: f"{x:g} yrs")
            table["Annual chance"] = table["return_period"].map(lambda x: f"{100 / x:.1f}%")
            table["Loss"] = table["loss_kes"].map(format_kes)
            st.dataframe(table[["Return period", "Annual chance", "Loss"]], hide_index=True, width="stretch")
        card_end()

    housing_col = first_existing(buildings_df.columns, ["housing_class", "class", "building_class"])
    loss_col = first_existing(buildings_df.columns, ["expected_annual_loss_kes", "annual_loss_kes", "loss_kes"])
    value_col = first_existing(buildings_df.columns, ["tiv_kes", "insured_value_kes", "value_kes"])

    bottom_left, bottom_right = st.columns([2, 1])
    with bottom_left:
        card_start("Loss by housing class", "API")
        if housing_col and loss_col and not buildings_df.empty:
            class_df = buildings_df.copy()
            class_df["_loss"] = class_df[loss_col].map(to_number)
            grouped = (
                class_df.dropna(subset=["_loss"])
                .groupby(housing_col, dropna=False)["_loss"]
                .sum()
                .sort_values(ascending=False)
                .reset_index()
            )
            fig = px.bar(
                grouped,
                x=housing_col,
                y="_loss",
                color=housing_col,
                color_discrete_sequence=CLASS_COLORS,
                labels={housing_col: "Class", "_loss": "Loss (KES)"},
                hover_data={"_loss": ":,.0f"},
            )
            fig.update_layout(showlegend=False, hovermode="closest")
            fig.update_xaxes(title="Housing class", tickangle=-20)
            fig.update_yaxes(
                title="Loss (KES)",
                tickformat="~s",
                rangemode="tozero",
                separatethousands=True,
            )
            st.plotly_chart(figure_layout(fig, height=360), width="stretch")
            st.caption(
                "Hover a bar for the class and exact loss total. Classes are sorted "
                "from highest to lowest loss."
            )
        else:
            st.info("Building class and loss fields are needed for this chart.")
        card_end()

    with bottom_right:
        card_start("Exposure by class")
        if housing_col and not buildings_df.empty:
            class_df = buildings_df.copy()
            if value_col:
                class_df["_value"] = class_df[value_col].map(to_number)
            if loss_col:
                class_df["_loss"] = class_df[loss_col].map(to_number)
            aggregations = {"_count": (housing_col, "count")}
            if value_col:
                aggregations["_value"] = ("_value", "sum")
            if loss_col:
                aggregations["_loss"] = ("_loss", "sum")
            grouped = class_df.groupby(housing_col).agg(**aggregations).reset_index()
            display = pd.DataFrame({"Class": grouped[housing_col], "Buildings": grouped["_count"]})
            if "_value" in grouped:
                display["Value"] = grouped["_value"].map(format_kes)
            if "_loss" in grouped:
                display["Loss"] = grouped["_loss"].map(format_kes)
            st.dataframe(display, hide_index=True, width="stretch")
        else:
            st.info("No housing class field was returned.")
        card_end()


def load_named_flood_hotspots():
    data_path = (
        Path(__file__).resolve().parent.parent
        / "data"
        / "nairobi_hotspots_geocoded.csv"
    )
    hotspots = pd.read_csv(data_path)
    required_columns = {"name", "lat", "lon"}
    if not required_columns.issubset(hotspots.columns):
        raise ValueError("Named hotspot data must contain name, lat, and lon columns.")
    hotspots = hotspots[["name", "lat", "lon"]].copy()
    hotspots["lat"] = pd.to_numeric(hotspots["lat"], errors="coerce")
    hotspots["lon"] = pd.to_numeric(hotspots["lon"], errors="coerce")
    hotspots = hotspots.dropna(subset=["name", "lat", "lon"])
    return hotspots.sort_values("name").reset_index(drop=True)


def render_hotspots():
    st.markdown("<h1>Flood hotspots</h1>", unsafe_allow_html=True)
    st.markdown(
        "<div class='hero-subtitle'>Explore the named Nairobi flood hotspots "
        "and their supplied geocoded locations.</div>",
        unsafe_allow_html=True,
    )
    st.markdown(
        "<div class='badge-row'><span class='badge real'>GEOCODED HOTSPOTS</span></div>",
        unsafe_allow_html=True,
    )

    hotspots = load_named_flood_hotspots()
    search = st.text_input(
        "Search hotspots",
        placeholder="Kiambiu, Dandora, Kibera...",
        key="hotspot_search",
    )
    filtered = hotspots
    if search:
        filtered = hotspots[
            hotspots["name"].astype(str).str.contains(search, case=False, na=False)
        ]

    map_col, list_col = st.columns([1.7, 1])
    with map_col:
        card_start("Named hotspot map", "24 LOCATIONS")
        if filtered.empty:
            st.info("No named hotspots match that search.")
        else:
            fig = px.scatter_map(
                filtered,
                lat="lat",
                lon="lon",
                hover_name="name",
                hover_data={"lat": ":.4f", "lon": ":.4f"},
                zoom=9,
                center={"lat": -1.2865, "lon": 36.8172},
                map_style="open-street-map",
                labels={"lat": "Latitude", "lon": "Longitude"},
            )
            fig.update_traces(
                marker=dict(size=14, color="#dc4b26", opacity=0.9),
                hovertemplate=(
                    "<b>%{hovertext}</b><br>"
                    "Latitude: %{lat:.4f}<br>"
                    "Longitude: %{lon:.4f}<extra></extra>"
                ),
            )
            fig.update_layout(
                height=570,
                margin=dict(l=0, r=0, t=0, b=0),
                hoverlabel=dict(
                    bgcolor="white",
                    bordercolor=LINE,
                    font=dict(color=TEXT),
                ),
            )
            st.plotly_chart(
                fig,
                width="stretch",
                config={"scrollZoom": True, "displaylogo": False},
            )
            st.caption(
                "Hover a marker to see the hotspot and coordinates. Use the map "
                "controls or scroll to zoom."
            )
        card_end()

    with list_col:
        card_start("Hotspot locations", "NAME · LAT · LON")
        if filtered.empty:
            st.info("No matching locations.")
        else:
            st.dataframe(
                filtered.rename(
                    columns={"name": "Name", "lat": "Latitude", "lon": "Longitude"}
                ),
                hide_index=True,
                width="stretch",
                height=570,
                column_config={
                    "Latitude": st.column_config.NumberColumn(format="%.4f"),
                    "Longitude": st.column_config.NumberColumn(format="%.4f"),
                },
            )
        st.caption(
            "Source: nairobi_hotspots_geocoded.csv. Coordinates are supplied "
            "reference locations, not live flood observations."
        )
        card_end()


def quote_payload_from_row(row):
    mapping = {
        "tiv_kes": ["tiv_kes", "insured_value_kes", "value_kes"],
        "expected_annual_loss_kes": ["expected_annual_loss_kes", "annual_loss_kes", "loss_kes"],
        "risk_class": ["risk_class", "risk_category"],
        "confidence": ["confidence", "model_confidence"],
    }
    payload = {}
    for target, candidates in mapping.items():
        for candidate in candidates:
            if candidate in row and pd.notna(row[candidate]):
                payload[target] = row[candidate]
                break
    return payload


def render_find_quote(token, buildings_df):
    st.markdown("<h1>Find a quote by location</h1>", unsafe_allow_html=True)
    st.markdown(
        "<div class='hero-subtitle'>Search for a Nairobi neighbourhood such as Kasarani or Lavington, compare matching assets, and request a quote.</div>",
        unsafe_allow_html=True,
    )

    if buildings_df.empty:
        st.info(        "No property records are currently available to search.")
        return

    building_col = first_existing(buildings_df.columns, ["building_id", "id"])
    loc_col = location_column(buildings_df)
    if not loc_col:
        st.warning(
            "Location search needs a neighbourhood, estate, ward, area, or address field in the property records."
        )
        st.dataframe(buildings_df.head(20), hide_index=True, width="stretch")
        return

    risk_col = first_existing(buildings_df.columns, ["risk_class", "risk_category"])
    loss_col = first_existing(buildings_df.columns, ["expected_annual_loss_kes", "annual_loss_kes", "loss_kes"])
    value_col = first_existing(buildings_df.columns, ["tiv_kes", "insured_value_kes", "value_kes"])

    search = st.text_input(
        "Search neighbourhood",
        placeholder="Try Kasarani or Lavington",
        help="Searches the location fields available in the property records.",
        key="quote_location_search",
    )
    matches = filter_location_records(buildings_df, search)
    location_options = sorted(
        matches[loc_col].dropna().astype(str).str.strip().loc[lambda values: values.ne("")]
        .unique()
        .tolist(),
        key=str.casefold,
    )
    if location_options:
        selected_location = st.selectbox(
            "Choose a matching neighbourhood",
            ["All matching locations", *location_options],
            key="quote_location_choice",
        )
        if selected_location != "All matching locations":
            matches = matches[
                matches[loc_col].astype("string").str.casefold()
                == selected_location.casefold()
            ].copy()

    left, right = st.columns([1.15, 1])
    with left:
        card_start("Matching assets", "API DATA")
        if matches.empty:
            st.info(f"No assets matched “{search.strip()}”. Try another neighbourhood name.")
            card_end()
            return
        cols = list(
            dict.fromkeys(
                col
                for col in [
                    building_col,
                    loc_col,
                    risk_col,
                    value_col,
                    loss_col,
                ]
                if col
            )
        )
        preview = matches[cols].copy() if cols else matches.copy()
        st.dataframe(preview.head(100), hide_index=True, width="stretch")

        metric_options = {}
        if value_col:
            metric_options["Insured value"] = value_col
        if loss_col:
            loss_metric_label = (
                "Reported loss"
                if loss_col.casefold() == "loss_kes"
                else "Expected annual loss"
            )
            metric_options[loss_metric_label] = loss_col
        risk_score_col = first_existing(
            matches.columns, ["risk_score", "aal_rate"]
        )
        if risk_score_col:
            metric_options["Risk score"] = risk_score_col

        if metric_options:
            chart_metric = st.selectbox(
                "Compare matching assets by",
                list(metric_options),
                key="quote_location_chart_metric",
            )
            chart_col = metric_options[chart_metric]
            chart_df = matches.copy()
            chart_df["_value"] = chart_df[chart_col].map(to_number)
            chart_df = chart_df.dropna(subset=["_value"])
            if building_col:
                chart_df["_asset"] = chart_df[building_col].astype(str)
                if chart_df["_asset"].duplicated().any():
                    chart_df["_asset"] = (
                        chart_df["_asset"]
                        + " · "
                        + chart_df[loc_col].astype(str)
                    )
            else:
                chart_df["_asset"] = chart_df[loc_col].astype(str)

            chart_df = chart_df.nlargest(15, "_value").sort_values("_value")
            if chart_df.empty:
                st.info(f"No numeric {chart_metric.lower()} values were returned.")
            else:
                fig = px.bar(
                    chart_df,
                    x="_value",
                    y="_asset",
                    orientation="h",
                    color="_value",
                    color_continuous_scale=["#b9e6e3", TEAL],
                    labels={"_value": chart_metric, "_asset": "Asset"},
                    hover_data={"_value": ":,.0f"},
                )
                fig.update_layout(coloraxis_showscale=False, showlegend=False)
                fig.update_yaxes(title=None)
                st.plotly_chart(
                    figure_layout(fig, height=350), width="stretch"
                )
                st.caption("This chart recalculates from the matching property records.")
        card_end()

    with right:
        card_start("Quote result")
        if building_col:
            matches = matches.reset_index(drop=True)

            def asset_label(index):
                row = matches.iloc[index]
                return f"{row[building_col]} · {row[loc_col]}"

            selected_index = st.selectbox(
                "Select an asset",
                range(len(matches)),
                format_func=asset_label,
                key="quote_selected_asset",
            )
            selected_row = matches.iloc[selected_index]

            if st.button("Get quote", type="primary", width="stretch"):
                with st.spinner("Fetching quote from backend..."):
                    try:
                        quote = get_building_quote(token, str(selected_row[building_col]))
                    except APIError as exc:
                        if exc.status_code != 404:
                            st.error(str(exc))
                            card_end()
                            return

                        payload = quote_payload_from_row(selected_row)
                        required = {"tiv_kes", "expected_annual_loss_kes", "risk_class", "confidence"}
                        if required.issubset(payload):
                            try:
                                quote = create_quote(token, payload)
                            except APIError as quote_exc:
                                st.error(str(quote_exc))
                            else:
                                st.dataframe(pd.json_normalize(quote), hide_index=True, width="stretch")
                        else:
                            st.error(
                                "The quote endpoint did not find this asset, and the returned asset fields are insufficient to request a quote."
                            )
                    else:
                        st.dataframe(pd.json_normalize(quote), hide_index=True, width="stretch")
        else:
            st.info(
                "The property records do not include an asset ID, so a quote cannot be requested."
            )
        card_end()


def render_drainage(buildings_df):
    st.markdown("<h1>Drainage reports</h1>", unsafe_allow_html=True)
    st.markdown("<div class='hero-subtitle'>A frontend view of backend-returned locations that may need drainage review.</div>", unsafe_allow_html=True)

    loc_col = location_column(buildings_df)
    risk_col = first_existing(buildings_df.columns, ["risk_class", "risk_category", "risk_score"])
    loss_col = first_existing(buildings_df.columns, ["expected_annual_loss_kes", "annual_loss_kes", "loss_kes"])

    if buildings_df.empty:
        st.info("No building/location data returned.")
        return

    card_start("How this improves the hazard layer")
    st.write("Use this page to quickly scan high-risk locations from the backend and identify places that may need drainage evidence or engineering review.")
    card_end()

    search = st.text_input("Search findings", placeholder="Search locations...")
    filtered = buildings_df.copy()
    if search and loc_col:
        filtered = filtered[filtered[loc_col].astype(str).str.contains(search, case=False, na=False)]

    if risk_col:
        filtered = filtered.sort_values(risk_col, ascending=False, key=lambda col: col.astype(str))
    elif loss_col:
        filtered = filtered.assign(_loss=filtered[loss_col].map(to_number)).sort_values("_loss", ascending=False)

    cols = st.columns(2)
    for index, (_, row) in enumerate(filtered.head(10).iterrows()):
        with cols[index % 2]:
            card_start(str(row.get(loc_col, row.get("building_id", "Location"))), "API")
            if risk_col:
                st.write(f"Risk: **{row.get(risk_col)}**")
            if loss_col:
                st.write(f"Expected loss: **{format_kes(row.get(loss_col))}**")
            available = {key: row[key] for key in row.index if pd.notna(row[key])}
            st.caption(", ".join(f"{prettify(k)}: {v}" for k, v in list(available.items())[:5]))
            card_end()


def render_methodology(formula, metrics, summary):
    st.markdown("<h1>Methodology & limitations</h1>", unsafe_allow_html=True)
    st.markdown("<div class='hero-subtitle'>Everything the model returns, and what it cannot tell you.</div>", unsafe_allow_html=True)

    card_start("Model chain")
    st.markdown(
        """
        <div style="display:flex;gap:14px;flex-wrap:wrap;align-items:center;">
            <span class="badge real" style="font-size:.95rem;padding:10px 14px;">Hazard</span>
            <b>→</b>
            <span class="badge real" style="font-size:.95rem;padding:10px 14px;">Vulnerability</span>
            <b>→</b>
            <span class="badge real" style="font-size:.95rem;padding:10px 14px;">Exposure</span>
            <b>→</b>
            <span class="badge real" style="font-size:.95rem;padding:10px 14px;">Financial engine</span>
            <b>→</b>
            <span class="badge real" style="font-size:.95rem;padding:10px 14px;">EP curve</span>
        </div>
        """,
        unsafe_allow_html=True,
    )
    card_end()

    card_start("Pricing formula", "API")
    if isinstance(formula, dict):
        formula_text = formula.get("pricing_formula") or formula.get("formula")
        if formula_text:
            st.code(str(formula_text))
        assumptions = formula.get("assumptions")
        if assumptions:
            st.write("Assumptions")
            if isinstance(assumptions, dict):
                st.dataframe(pd.json_normalize(assumptions), hide_index=True, width="stretch")
            else:
                st.write(assumptions)
    else:
        st.info("The formula endpoint did not return a dictionary payload.")
    card_end()

    left, right = st.columns(2)
    with left:
        card_start("Model metrics", "API")
        if isinstance(metrics, dict) and metrics:
            st.dataframe(pd.json_normalize(metrics), hide_index=True, width="stretch")
        else:
            st.info("No metrics payload returned.")
        card_end()

    with right:
        card_start("Portfolio summary payload", "API")
        if isinstance(summary, dict) and summary:
            st.dataframe(pd.json_normalize(summary), hide_index=True, width="stretch")
        else:
            st.info("No summary payload returned.")
        card_end()


def render_report(summary, buildings_df, rp_df):
    st.markdown("<h1>Generate a report</h1>", unsafe_allow_html=True)
    st.markdown("<div class='hero-subtitle'>Create a plain-English briefing from the latest backend data shown in this dashboard.</div>", unsafe_allow_html=True)

    title = st.text_input("Report title", value="Nairobi Pluvial Flood Risk Briefing")
    audience = st.segmented_control(
        "Audience",
        ["Underwriter", "Portfolio manager", "County / disaster body", "Broker / cedant"],
        default="Underwriter",
    )
    include_sections = st.multiselect(
        "Sections",
        ["EP curve", "Class breakdown", "Hotspots", "Quote guidance", "Assumptions & limitations"],
        default=["EP curve", "Class breakdown", "Hotspots", "Assumptions & limitations"],
    )

    total_assets = len(buildings_df)
    report = [
        f"# {title}",
        "",
        f"Audience: {audience}",
        "",
        f"The current backend response contains {total_assets} building records.",
    ]
    if not rp_df.empty:
        worst = rp_df.sort_values("loss_kes").iloc[-1]
        report.append(f"The largest return-period loss currently returned is {format_kes(worst['loss_kes'])} at {worst['return_period']:.0f} years.")
    if "Class breakdown" in include_sections:
        housing_col = first_existing(buildings_df.columns, ["housing_class", "class", "building_class"])
        if housing_col:
            report.append(f"The portfolio includes {buildings_df[housing_col].nunique()} returned housing classes.")
    report.append("")
    report.append("This report is generated in the Streamlit frontend from authenticated backend responses.")
    report_text = "\n".join(report)

    left, right = st.columns([1, 1.5])
    with left:
        st.download_button("Download report", report_text, file_name="nairobi-flood-risk-briefing.md", width="stretch")
    with right:
        card_start("Report preview")
        st.markdown(report_text)
        card_end()


def render_bulk_upload():
    st.markdown("<h1>Bulk upload</h1>", unsafe_allow_html=True)
    st.markdown(
        "<div class='hero-subtitle'>Upload supported documents to the underwriting "
        "service and generate a downloadable local summary report.</div>",
        unsafe_allow_html=True,
    )
    st.caption(
        "The local report extracts text and previews structured data; it is not an "
        "AI-generated underwriting review or a flood-risk calculation."
    )
    st.caption(
        "Supported: .txt, .md, .docx, .pdf, .xlsx, .csv, .html, .json, .py, "
        ".png, .jpg/.jpeg, .svg, .tif/.tiff (GeoTIFF), and .zip."
    )
    st.info(
        "Python files are read as text and never run. Raster images and GeoTIFFs "
        "provide metadata only (no OCR or flood-risk calculation); SVG text labels "
        "are extracted. Scanned PDFs may have no extractable text."
    )

    uploads = st.file_uploader(
        "Choose one or more files",
        type=list(UPLOAD_EXTENSIONS),
        accept_multiple_files=True,
        key="bulk_upload_input",
        help="Up to 25 MB per file and 50 MB total. ZIP contents are inspected without being extracted to disk.",
    )
    if not uploads:
        st.session_state.pop("bulk_upload_result", None)
        st.caption("Select files, then choose Generate summary.")
        return

    file_data = [(uploaded.name, uploaded.getvalue()) for uploaded in uploads]
    total_bytes = sum(len(content) for _, content in file_data)
    fingerprint = sha256(
        b"".join(
            name.encode("utf-8") + b"\0" + sha256(content).digest()
            for name, content in file_data
        )
    ).hexdigest()
    if total_bytes > MAX_UPLOAD_TOTAL_BYTES:
        st.error("The selected files exceed the 50 MB total upload limit.")
        st.session_state.pop("bulk_upload_result", None)
        return

    st.dataframe(
        pd.DataFrame(
            [
                {"File": name, "Size": f"{len(content) / 1024:,.1f} KB"}
                for name, content in file_data
            ]
        ),
        hide_index=True,
        width="stretch",
    )

    st.markdown("#### Submit a document to the backend")
    st.caption(
        "Choose a portfolio name and document to send one file to the underwriting "
        "service for extraction and analysis."
    )
    selected_upload = st.selectbox(
        "Document to upload",
        range(len(file_data)),
        format_func=lambda index: file_data[index][0],
        key="bulk_upload_backend_file",
    )
    portfolio_name = st.text_input(
        "Portfolio name",
        value="Commercial Property Portfolio",
        key="bulk_upload_portfolio_name",
    ).strip()
    selected_name, selected_content = file_data[selected_upload]
    upload_fingerprint = sha256(
        f"{fingerprint}:{selected_upload}:{portfolio_name}".encode("utf-8")
    ).hexdigest()
    if (
        st.session_state.get("bulk_upload_document_fingerprint")
        != upload_fingerprint
    ):
        for key in (
            "bulk_upload_document_summary",
            "bulk_upload_document_summary_error",
            "bulk_upload_document_chat",
            "bulk_upload_document_chat_error",
        ):
            st.session_state.pop(key, None)
        st.session_state["bulk_upload_document_fingerprint"] = upload_fingerprint

    backend_result = st.session_state.get("bulk_upload_backend_result")
    if backend_result and backend_result["fingerprint"] != upload_fingerprint:
        st.session_state.pop("bulk_upload_backend_result", None)
        backend_result = None

    if st.button(
        "Upload selected document to backend",
        type="primary",
        width="stretch",
        disabled=not portfolio_name or len(selected_content) > MAX_UPLOAD_BYTES,
        key="bulk_upload_backend_submit",
    ):
        for key in (
            "bulk_upload_document_summary",
            "bulk_upload_document_summary_error",
            "bulk_upload_document_chat",
            "bulk_upload_document_chat_error",
        ):
            st.session_state.pop(key, None)
        try:
            with st.spinner("Uploading document and waiting for extraction..."):
                response = upload_document(
                    TOKEN,
                    selected_name,
                    selected_content,
                    portfolio_name,
                )
            st.session_state["bulk_upload_backend_result"] = {
                "fingerprint": upload_fingerprint,
                "response": response,
                "error": None,
            }
        except APIError as exc:
            st.session_state["bulk_upload_backend_result"] = {
                "fingerprint": upload_fingerprint,
                "response": None,
                "error": str(exc),
                "status_code": exc.status_code,
            }
        backend_result = st.session_state.get("bulk_upload_backend_result")

    if len(selected_content) > MAX_UPLOAD_BYTES:
        st.warning("This document exceeds the 25 MB per-file upload limit.")
    if backend_result and backend_result["fingerprint"] == upload_fingerprint:
        if backend_result["error"]:
            status = (
                f" (HTTP {backend_result['status_code']})"
                if backend_result.get("status_code") is not None
                else ""
            )
            st.error(f"Backend upload failed{status}: {backend_result['error']}")
        elif backend_result["response"] is None:
            st.info("The backend completed the request without a response body.")
        else:
            response = backend_result["response"]
            document_payload = (
                response if isinstance(response, dict) else {"response": response}
            )
            readiness = (
                document_payload.get("model_readiness")
                if isinstance(response, dict)
                else None
            )
            if isinstance(readiness, dict) and readiness.get("ready") is False:
                st.warning(
                    "The document was received, but the backend reports that the "
                    "extracted records are not ready for a model run. Review the "
                    "issues and warnings in the response below."
                )
            else:
                st.success("Document uploaded and processed by the backend.")
            st.json(response, expanded=True)

            st.markdown("#### AI summary and chat with this document")
            st.caption(
                "The authenticated AI service uses the uploaded document "
                "and extracted records as context. Its output is explanatory only."
            )
            portfolio_summary = document_payload.get("portfolio_summary")
            if not isinstance(portfolio_summary, dict):
                portfolio_summary = document_payload.get("financial_summary")
            if not isinstance(portfolio_summary, dict):
                portfolio_summary = {}
            portfolio_summary = {
                **portfolio_summary,
                "portfolio_name": portfolio_name,
                "document_type": document_payload.get("document_type"),
                "document_summary": document_payload.get("summary"),
            }
            building_risk_summary = document_payload.get(
                "building_risk_summary"
            )
            if not isinstance(building_risk_summary, list):
                building_risk_summary = document_payload.get(
                    "extracted_rows", []
                )
            if not isinstance(building_risk_summary, list):
                building_risk_summary = []
            document_context = {
                "portfolio_summary": portfolio_summary,
                "building_risk_summary": building_risk_summary,
                "document": document_payload,
                "portfolio_name": portfolio_name,
            }
            if st.button(
                "Generate AI summary",
                type="primary",
                key="bulk_upload_generate_ai_summary",
            ):
                try:
                    with st.spinner("Generating an underwriter-friendly summary..."):
                        summary_response = request_uploaded_document_summary(
                            TOKEN,
                            portfolio_summary,
                            building_risk_summary,
                        )
                    st.session_state["bulk_upload_document_summary"] = (
                        summary_response
                    )
                    st.session_state.pop(
                        "bulk_upload_document_summary_error", None
                    )
                except APIError as exc:
                    st.session_state["bulk_upload_document_summary_error"] = str(exc)
                    st.session_state.pop("bulk_upload_document_summary", None)

            summary_error = st.session_state.get(
                "bulk_upload_document_summary_error"
            )
            if summary_error:
                st.error(f"AI summary failed: {summary_error}")
            elif "bulk_upload_document_summary" in st.session_state:
                summary_response = st.session_state[
                    "bulk_upload_document_summary"
                ]
                summary_text = interpretation_response_text(
                    summary_response,
                    ("summary", "underwriter_explanation", "message"),
                )
                if summary_text:
                    card_start("AI-generated document summary", "INTERPRETATION ONLY")
                    st.markdown(summary_text)
                    if isinstance(summary_response, dict):
                        source = summary_response.get("source")
                        if source:
                            st.caption(f"Source: {safe(source)}")
                    card_end()
                else:
                    st.info(
                        "The AI service returned a response without a "
                        "summary field."
                    )
                    st.json(summary_response, expanded=True)

            chat_history = st.session_state.get(
                "bulk_upload_document_chat", []
            )
            chat_question = st.chat_input(
                "Ask a question about this uploaded document...",
                key="bulk_upload_document_chat_input",
            )
            if chat_question:
                chat_history.append(
                    {"role": "user", "content": chat_question}
                )
                try:
                    with st.spinner("Preparing an answer from the document..."):
                        answer_response = ask_uploaded_document_question(
                            TOKEN,
                            document_context,
                            chat_question,
                        )
                    answer_text = interpretation_response_text(
                        answer_response,
                        ("answer", "response", "summary", "message"),
                    )
                    chat_history.append(
                        {
                            "role": "assistant",
                            "content": answer_text
                            or "The AI service returned no answer text.",
                            "source": (
                                answer_response.get("source")
                                if isinstance(answer_response, dict)
                                else None
                            ),
                        }
                    )
                    st.session_state.pop(
                        "bulk_upload_document_chat_error", None
                    )
                except APIError as exc:
                    st.session_state["bulk_upload_document_chat_error"] = str(exc)
                st.session_state["bulk_upload_document_chat"] = chat_history

            for message in st.session_state.get(
                "bulk_upload_document_chat", []
            ):
                with st.chat_message(message["role"]):
                    st.markdown(message["content"])
                    if message.get("source"):
                        st.caption(f"Source: {safe(message['source'])}")
            chat_error = st.session_state.get(
                "bulk_upload_document_chat_error"
            )
            if chat_error:
                st.error(f"Document chat failed: {chat_error}")

    st.markdown("#### Generate a local summary")
    st.caption(
        "This separate action summarizes selected files in the dashboard and does "
        "not send them to the backend."
    )

    previous_result = st.session_state.get("bulk_upload_result")
    if previous_result and previous_result["fingerprint"] != fingerprint:
        st.session_state.pop("bulk_upload_result", None)
        previous_result = None

    if st.button("Generate summary", type="primary", width="stretch"):
        results = []
        report_sections = [
            "# Bulk upload summary",
            "",
            f"Files submitted: {len(file_data)}",
            "",
            "This report summarizes uploaded files in the Streamlit frontend. "
            "It does not calculate flood risk or send files to the backend.",
        ]
        for name, content in file_data:
            if len(content) > MAX_UPLOAD_BYTES:
                status, details = "Error", "File exceeds the 25 MB per-file processing limit."
            else:
                try:
                    details = summarize_upload_content(name, content)
                    status = "Processed"
                except (
                    ValueError, OSError, UnicodeError, PdfReadError,
                    InvalidFileException, Image.DecompressionBombError,
                    zipfile.BadZipFile,
                ) as exc:
                    status, details = "Error", str(exc)
            results.append(
                {
                    "File": name,
                    "Status": status,
                    "Summary": details,
                }
            )
            report_sections.extend(
                [
                    "",
                    f"## {name} — {status}",
                    "",
                    details,
                ]
            )

        st.session_state["bulk_upload_result"] = {
            "fingerprint": fingerprint,
            "results": results,
            "report": "\n".join(report_sections),
        }
        previous_result = st.session_state["bulk_upload_result"]

    if previous_result:
        results = previous_result["results"]
        error_count = sum(result["Status"] == "Error" for result in results)
        if error_count:
            st.warning(
                f"Summary generated with {error_count} file(s) that could not be read. "
                "See each file's details below."
            )
        else:
            st.success(f"Summary generated for {len(results)} file(s).")

        st.download_button(
            "Download generated summary (.md)",
            previous_result["report"],
            file_name="bulk-upload-summary.md",
            mime="text/markdown",
            width="stretch",
        )
        for result in results:
            with st.expander(f"{result['Status']}: {result['File']}"):
                if result["Status"] == "Error":
                    st.error(result["Summary"])
                else:
                    st.code(result["Summary"], language="text")


def render_api_explorer(token):
    st.markdown("<h1>🔌 Data explorer</h1>", unsafe_allow_html=True)
    st.markdown(
        "<div class='hero-subtitle'>Inspect the available model-data connections. "
        "Requests here are separate from Bulk Upload.</div>",
        unsafe_allow_html=True,
    )
    st.info(
        "This explorer sends read-only GET requests using your signed-in session. "
        "It does not create or modify backend records."
    )

    endpoint_labels = {
        "Model health": "health",
        "Buildings": "buildings",
        "Portfolio summary": "portfolio-summary",
        "Metrics": "metrics",
        "Pricing formula": "formula",
        "Quote for building": None,
    }
    selected_label = st.selectbox(
        "Endpoint",
        list(endpoint_labels),
        key="api_explorer_endpoint",
    )
    building_id = ""
    if endpoint_labels[selected_label] is None:
        building_id = st.text_input(
            "Building ID",
            key="api_explorer_building_id",
            help="The identifier from a building record returned by the Buildings endpoint.",
        ).strip()

    params_text = st.text_area(
        "Query parameters (JSON)",
        value="{}",
        key="api_explorer_params",
        help="Optional query parameters, for example {\"search\": \"Kasarani\"}.",
    )
    endpoint = endpoint_labels[selected_label]
    if endpoint is None:
        endpoint = f"quote/{quote(building_id, safe='')}" if building_id else "quote/{building_id}"
    route = f"/api/model/{endpoint}/"
    st.code(f"GET {route}", language="http")

    if st.button(
        "Send GET request",
        type="primary",
        disabled=selected_label == "Quote for building" and not building_id,
        key="api_explorer_send",
    ):
        try:
            params = json.loads(params_text)
            if not isinstance(params, dict):
                raise ValueError("Query parameters must be a JSON object.")
        except json.JSONDecodeError as exc:
            st.error(f"Invalid query-parameter JSON: {exc}")
            st.session_state.pop("api_explorer_result", None)
        except ValueError as exc:
            st.error(str(exc))
            st.session_state.pop("api_explorer_result", None)
        else:
            try:
                with st.spinner("Requesting the backend..."):
                    response = model_get(endpoint, token, params=params or None)
                st.session_state["api_explorer_result"] = {
                    "route": route,
                    "response": response,
                    "error": None,
                }
            except APIError as exc:
                st.session_state["api_explorer_result"] = {
                    "route": route,
                    "response": None,
                    "error": str(exc),
                    "status_code": exc.status_code,
                }

    result = st.session_state.get("api_explorer_result")
    if result and result["route"] == route:
        st.markdown("#### Response")
        if result["error"]:
            status = (
                f" (HTTP {result['status_code']})"
                if result.get("status_code") is not None
                else ""
            )
            st.error(f"Request failed{status}: {result['error']}")
        elif result["response"] is None:
            st.info("The endpoint returned an empty response.")
        else:
            st.json(result["response"], expanded=True)


if not TOKEN:
    inject_styles()
    st.warning("Please sign in to view the dashboard.")
    if st.button("Go to login"):
        st.switch_page("views/login.py")
    st.stop()

inject_styles()
page = render_sidebar()

if page in {"BULK UPLOAD", "Data Explorer", "Hotspots"}:
    render_dashboard_header(TOKEN, pd.DataFrame())
    if page == "BULK UPLOAD":
        render_bulk_upload()
    elif page == "Data Explorer":
        render_api_explorer(TOKEN)
    else:
        render_hotspots()
else:
    try:
        with st.spinner("Loading flood-risk data..."):
            data = load_dashboard_data(TOKEN)
    except AuthenticationError:
        st.error("Your session has expired or is invalid. Please sign in again.")
        logout()
    except APIError as exc:
        st.error(str(exc))
        st.stop()

    health = data["health"]
    if isinstance(health, dict) and health.get("status"):
        st.caption(f"☁️ Service status: {health.get('status')}")

    summary = data["summary"] or {}
    metrics = data["metrics"] or {}
    formula = data["formula"] or {}
    buildings_df = pd.DataFrame(normalize_records(data["buildings"]))
    rp_df = return_period_df(summary)

    render_dashboard_header(TOKEN, buildings_df)

    if page == "Overview":
        render_overview(summary, metrics, formula, buildings_df, rp_df)
    elif page == "EP Curve":
        render_ep_curve(summary, buildings_df, rp_df)
    elif page == "Drainage Reports":
        render_drainage(buildings_df)
    elif page == "Find Quote":
        render_find_quote(TOKEN, buildings_df)
    elif page == "Methodology":
        render_methodology(formula, metrics, summary)
    elif page == "Interpretation":
        render_interpretation_page(
            summary, metrics, formula, buildings_df
        )
    else:
        render_report(summary, buildings_df, rp_df)
